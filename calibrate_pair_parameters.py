#!/usr/bin/env python3
"""Preview or adopt a coherent pair-centroid and timing calibration.

Preview mode computes, without modifying the canonical workbook:

1. a direct-P common centroid for each configured pair,
2. a pair-relative ``new time shift``, and
3. a common ``pick align shift`` after the relative alignment.

Adoption is a separate, guarded action.  It writes all five workbook fields
(``new_lat``, ``new_lon``, ``new_depth``, ``new time shift``, and
``pick align shift``) together for selected pairs only when every selected
candidate is complete and the workbook still matches the preview input hash.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import shutil
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl

import make_multiphase_median_outputs as multiphase
import relocate_pair_fixed_depths as centroid


ROOT = Path(__file__).resolve().parent
DEFAULT_CONFIG = ROOT / "analysis_config.json"
CANDIDATE_FILENAME = "pair_calibration_candidates.csv"
MANIFEST_FILENAME = "pair_calibration_manifest.json"
ADOPTED_FIELDS = (
    "new_lat",
    "new_lon",
    "new_depth",
    "new time shift",
    "pick align shift",
)


class CalibrationError(RuntimeError):
    """Raised when a calibration preview or adoption is unsafe."""


def finite_float(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_header(values: tuple[Any, ...]) -> dict[str, int]:
    return {
        str(value).strip().lower(): index + 1
        for index, value in enumerate(values)
        if value is not None and str(value).strip()
    }


def read_csv_rows(path: Path, key: str) -> dict[str, dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            label = str(row.get(key, "")).strip()
            if label:
                rows[label] = row
    return rows


def read_current_pair_values(
    workbook_path: Path, labels: list[str]
) -> dict[str, dict[str, float | None]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["pairs"]
        rows = sheet.iter_rows(values_only=True)
        header = normalized_header(next(rows))
        required = {"label", *ADOPTED_FIELDS}
        missing = sorted(required.difference(header))
        if missing:
            raise CalibrationError(
                "Workbook pairs sheet is missing columns: " + ", ".join(missing)
            )
        selected = set(labels)
        values: dict[str, dict[str, float | None]] = {}
        for row in rows:
            raw_label = row[header["label"] - 1]
            if raw_label is None:
                continue
            label = str(raw_label).strip()
            if label not in selected:
                continue
            values[label] = {
                field: finite_float(row[header[field] - 1]) for field in ADOPTED_FIELDS
            }
    finally:
        workbook.close()
    missing_labels = [label for label in labels if label not in values]
    if missing_labels:
        raise CalibrationError(
            "Configured pairs missing from workbook: " + ", ".join(missing_labels)
        )
    return values


def prepare_candidate_workbook(
    source_path: Path,
    destination_path: Path,
    labels: list[str],
    centroid_rows: dict[str, dict[str, str]],
) -> None:
    """Copy the workbook and install only successful candidate centroids."""
    shutil.copy2(source_path, destination_path)
    workbook = openpyxl.load_workbook(destination_path)
    try:
        sheet = workbook["pairs"]
        header = normalized_header(
            tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        )
        required = {"label", "new_lat", "new_lon", "new_depth"}
        missing = sorted(required.difference(header))
        if missing:
            raise CalibrationError(
                "Workbook pairs sheet is missing columns: " + ", ".join(missing)
            )
        selected = set(labels)
        found: set[str] = set()
        for row_index in range(2, sheet.max_row + 1):
            raw_label = sheet.cell(row_index, header["label"]).value
            if raw_label is None:
                continue
            label = str(raw_label).strip()
            if label not in selected:
                continue
            found.add(label)
            # Clear old overrides in the disposable copy first.  A failed
            # centroid fit must not silently feed a stale location into the
            # timing preview.
            for field in ("new_lat", "new_lon", "new_depth"):
                sheet.cell(row_index, header[field]).value = None
            row = centroid_rows.get(label, {})
            if str(row.get("status", "")) != "ok":
                continue
            candidates = {
                "new_lat": finite_float(row.get("new_lat")),
                "new_lon": finite_float(row.get("new_lon")),
                "new_depth": finite_float(row.get("new_depth")),
            }
            if any(value is None for value in candidates.values()):
                continue
            for field, value in candidates.items():
                cell = sheet.cell(row_index, header[field])
                cell.value = value
                cell.number_format = "0.0000" if field != "new_depth" else "0.0"
        missing_labels = [label for label in labels if label not in found]
        if missing_labels:
            raise CalibrationError(
                "Configured pairs missing from workbook: " + ", ".join(missing_labels)
            )
        workbook.save(destination_path)
    finally:
        workbook.close()


def difference(candidate: float | None, current: float | None) -> float | None:
    if candidate is None or current is None:
        return None
    return candidate - current


def build_candidate_rows(
    labels: list[str],
    current: dict[str, dict[str, float | None]],
    centroid_rows: dict[str, dict[str, str]],
    timing_rows: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for label in labels:
        centroid_row = centroid_rows.get(label, {})
        timing_row = timing_rows.get(label, {})
        candidates = {
            "new_lat": finite_float(centroid_row.get("new_lat")),
            "new_lon": finite_float(centroid_row.get("new_lon")),
            "new_depth": finite_float(centroid_row.get("new_depth")),
            "new time shift": finite_float(
                timing_row.get("computed_median_shift_seconds")
            ),
            "pick align shift": finite_float(
                timing_row.get("common_time_shift_seconds")
            ),
        }
        centroid_ok = str(centroid_row.get("status", "")) == "ok"
        timing_ok = all(
            candidates[field] is not None
            for field in ("new time shift", "pick align shift")
        )
        ready = centroid_ok and timing_ok and all(
            candidates[field] is not None for field in ADOPTED_FIELDS
        )
        row: dict[str, Any] = {
            "pair_label": label,
            "event1": timing_row.get("event1", centroid_row.get("event1", "")),
            "event2": timing_row.get("event2", centroid_row.get("event2", "")),
            "centroid_status": centroid_row.get("status", "missing"),
            "centroid_accepted_picks": centroid_row.get("accepted_picks", ""),
            "timing_good_count": timing_row.get("good_count", ""),
            "pick_align_high_snr_p_count": timing_row.get(
                "origin_fit_p_pick_count", ""
            ),
            "ready_to_adopt": ready,
        }
        for field in ADOPTED_FIELDS:
            key = field.replace(" ", "_")
            row[f"current_{key}"] = current[label][field]
            row[f"candidate_{key}"] = candidates[field]
            row[f"change_{key}"] = difference(candidates[field], current[label][field])
        rows.append(row)
    return rows


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def preview(config_path: Path, output: Path | None = None) -> Path:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    mode = centroid.centroid_location_mode(config)
    if mode == "catalog_fixed":
        raise CalibrationError(
            "centroid_location_mode=catalog_fixed disables centroid relocation"
        )
    labels = [str(label) for label in config["pairs"]]
    canonical_workbook = Path(
        config.get("time_shift_workbook") or config["catalog_path"]
    ).resolve()
    initial_hash = sha256(canonical_workbook)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = output or ROOT / "outputs" / f"pair_calibration_preview_{stamp}"
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=False)

    centroid_output = output / "centroid_relocation"
    print("STEP 1/3: computing candidate common centroids", flush=True)
    centroid.run(config_path, centroid_output, labels)
    centroid_rows = read_csv_rows(
        centroid_output / "pair_fixed_depth_location_summary.csv", "pair_label"
    )

    candidate_workbook = output / "candidate_ICevents_full.xlsx"
    prepare_candidate_workbook(
        canonical_workbook, candidate_workbook, labels, centroid_rows
    )
    candidate_config = dict(config)
    candidate_config.update(
        {
            "catalog_path": str(candidate_workbook),
            "time_shift_workbook": str(candidate_workbook),
            "time_shift_source": "computed",
            "common_time_shift_source": "high_snr_picks",
            "update_workbook_time_shifts_from_computed": False,
            "do_centroid_relocation": False,
        }
    )
    candidate_config_path = output / "candidate_analysis_config.json"
    candidate_config_path.write_text(
        json.dumps(candidate_config, indent=2) + "\n", encoding="utf-8"
    )

    print("STEP 2/3: computing candidate relative and common timing shifts", flush=True)
    multiphase_output = multiphase.run(candidate_config_path)
    timing_rows = read_csv_rows(multiphase_output / "median_summary.csv", "pair_label")
    current = read_current_pair_values(canonical_workbook, labels)
    candidate_rows = build_candidate_rows(
        labels, current, centroid_rows, timing_rows
    )
    candidate_csv = output / CANDIDATE_FILENAME
    write_csv(candidate_csv, candidate_rows)

    final_hash = sha256(canonical_workbook)
    if final_hash != initial_hash:
        raise CalibrationError(
            "Canonical workbook changed during preview; candidates are not safe to adopt"
        )
    ready_count = sum(str(row["ready_to_adopt"]).lower() == "true" for row in candidate_rows)
    manifest = {
        "created_utc": datetime.now(UTC).isoformat(timespec="seconds").replace(
            "+00:00", "Z"
        ),
        "mode": "preview_only_no_canonical_workbook_writes",
        "config_path": str(config_path),
        "canonical_workbook": str(canonical_workbook),
        "canonical_workbook_sha256": initial_hash,
        "candidate_workbook": str(candidate_workbook),
        "centroid_output": str(centroid_output),
        "multiphase_output": str(multiphase_output.resolve()),
        "candidate_csv": str(candidate_csv),
        "configured_pairs": labels,
        "ready_to_adopt_count": ready_count,
        "blocked_pair_count": len(labels) - ready_count,
        "manual_pick_alignment_workbook": candidate_config.get(
            "manual_pick_alignment_workbook", ""
        ),
    }
    (output / MANIFEST_FILENAME).write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    print("STEP 3/3: review table written; canonical workbook unchanged", flush=True)
    print(f"ready_to_adopt={ready_count}/{len(labels)}", flush=True)
    print(output, flush=True)
    return output


def parse_candidate_bool(value: Any) -> bool:
    return str(value).strip().lower() == "true"


def adopt(candidate_csv: Path, selected_pairs: list[str] | None = None) -> Path:
    candidate_csv = candidate_csv.resolve()
    manifest_path = candidate_csv.parent / MANIFEST_FILENAME
    if not manifest_path.is_file():
        raise CalibrationError(f"Missing calibration manifest: {manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    workbook_path = Path(manifest["canonical_workbook"]).resolve()
    expected_hash = str(manifest["canonical_workbook_sha256"])
    if sha256(workbook_path) != expected_hash:
        raise CalibrationError(
            "Canonical workbook has changed since preview; generate a fresh preview"
        )
    rows = read_csv_rows(candidate_csv, "pair_label")
    labels = selected_pairs or [str(label) for label in manifest["configured_pairs"]]
    missing = [label for label in labels if label not in rows]
    if missing:
        raise CalibrationError("Candidate rows missing for: " + ", ".join(missing))
    blocked = [
        label for label in labels if not parse_candidate_bool(rows[label]["ready_to_adopt"])
    ]
    if blocked:
        raise CalibrationError(
            "Adoption blocked because candidates are incomplete for: "
            + ", ".join(blocked)
        )
    parsed: dict[str, dict[str, float]] = {}
    for label in labels:
        parsed[label] = {}
        for field in ADOPTED_FIELDS:
            key = f"candidate_{field.replace(' ', '_')}"
            value = finite_float(rows[label].get(key))
            if value is None:
                raise CalibrationError(f"{label} has no finite {field!r} candidate")
            parsed[label][field] = value

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = workbook_path.with_name(
        f"{workbook_path.stem}_before_pair_calibration_{stamp}{workbook_path.suffix}"
    )
    shutil.copy2(workbook_path, backup)
    workbook = openpyxl.load_workbook(workbook_path)
    temporary = workbook_path.with_name(f".{workbook_path.name}.calibration.tmp")
    try:
        sheet = workbook["pairs"]
        header = normalized_header(
            tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        )
        required = {"label", *ADOPTED_FIELDS}
        missing_columns = sorted(required.difference(header))
        if missing_columns:
            raise CalibrationError(
                "Workbook pairs sheet is missing columns: "
                + ", ".join(missing_columns)
            )
        written: set[str] = set()
        for row_index in range(2, sheet.max_row + 1):
            raw_label = sheet.cell(row_index, header["label"]).value
            if raw_label is None:
                continue
            label = str(raw_label).strip()
            if label not in parsed:
                continue
            for field, value in parsed[label].items():
                cell = sheet.cell(row_index, header[field])
                cell.value = value
                cell.number_format = "0.0" if field == "new_depth" else "0.0000"
            written.add(label)
        missing_rows = [label for label in labels if label not in written]
        if missing_rows:
            raise CalibrationError(
                "Workbook rows missing during adoption: " + ", ".join(missing_rows)
            )
        workbook.save(temporary)
    finally:
        workbook.close()
    os.replace(temporary, workbook_path)
    adoption_record = {
        "adopted_utc": datetime.now(UTC).isoformat(timespec="seconds").replace(
            "+00:00", "Z"
        ),
        "candidate_csv": str(candidate_csv),
        "canonical_workbook": str(workbook_path),
        "backup": str(backup),
        "pairs": labels,
        "fields": list(ADOPTED_FIELDS),
        "resulting_workbook_sha256": sha256(workbook_path),
    }
    record_path = candidate_csv.parent / "pair_calibration_adoption.json"
    record_path.write_text(json.dumps(adoption_record, indent=2) + "\n")
    print(f"backup={backup}")
    print(f"written_pairs={len(labels)}")
    print(record_path)
    return record_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument(
        "--adopt-candidates",
        type=Path,
        default=None,
        help="Explicitly adopt a reviewed pair_calibration_candidates.csv file.",
    )
    parser.add_argument(
        "--pairs",
        nargs="*",
        default=None,
        help="With --adopt-candidates, adopt only these complete pair rows.",
    )
    args = parser.parse_args()
    os.chdir(ROOT)
    if args.adopt_candidates is not None:
        adopt(args.adopt_candidates, args.pairs)
    else:
        if args.pairs:
            raise SystemExit("--pairs is only valid with --adopt-candidates")
        preview(args.config.resolve(), args.output)


if __name__ == "__main__":
    main()
