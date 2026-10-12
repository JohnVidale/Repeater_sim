#!/usr/bin/env python3
"""Apply the display-unit contract for the working repeater workbook."""
from __future__ import annotations

import argparse
import math
import shutil
from datetime import datetime, timezone
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

TIME = "yyyy-mm-dd hh:mm:ss.000"
COORD = "0.00000"
METRES = "0"
SECONDS = "0.000"

def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--workbook", type=Path, default=Path("ICevents_working.xlsx"))
    args = p.parse_args()
    backup = args.workbook.parent / "archive" / f"{args.workbook.stem}_before_display_units_{datetime.now():%Y%m%d_%H%M%S}{args.workbook.suffix}"
    backup.parent.mkdir(exist_ok=True); shutil.copy2(args.workbook, backup)
    wb = load_workbook(args.workbook)
    try:
        aligned = wb["aligned_pairs"]
        h = {str(c.value).strip(): c.column for c in aligned[1] if c.value is not None}
        # A differential location is a distance, not a second pair of angular
        # coordinates.  Store its north/east components in whole metres.
        lat_col = h["same_location_event1_lat"]
        north_col = h.get("relative_location_delta_lat_degrees")
        east_col = h.get("relative_location_delta_lon_degrees")
        if north_col is not None:
            aligned.cell(1, north_col).value = "relative_location_north_m"
            aligned.cell(1, east_col).value = "relative_location_east_m"
            for row in range(2, aligned.max_row + 1):
                north, east, lat = aligned.cell(row, north_col).value, aligned.cell(row, east_col).value, aligned.cell(row, lat_col).value
                aligned.cell(row, north_col).value = None if north is None else 111195.0 * float(north)
                aligned.cell(row, east_col).value = None if east is None or lat is None else 111195.0 * math.cos(math.radians(float(lat))) * float(east)
        separation_col = h.get("relative_location_3d_km")
        if separation_col is not None:
            aligned.cell(1, separation_col).value = "relative_location_3d_m"
            for row in range(2, aligned.max_row + 1):
                value = aligned.cell(row, separation_col).value
                aligned.cell(row, separation_col).value = None if value is None else 1000.0 * float(value)
        for sheet in wb.worksheets:
            for cell in sheet[1]:
                if cell.value is None: continue
                name = str(cell.value).strip().lower()
                col = cell.column
                values = [sheet.cell(row, col).value for row in range(2, sheet.max_row + 1)]
                if name in {"relative_location_north_m", "relative_location_east_m", "relative_location_3d_m"}:
                    fmt, width = METRES, 10
                elif name == "time":
                    for row, value in enumerate(values, start=2):
                        if isinstance(value, str) and value.strip():
                            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
                            sheet.cell(row, col).value = parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed
                    fmt, width = TIME, 23
                elif "time" in name or "timestamp" in name:
                    fmt = TIME if any(isinstance(value, datetime) for value in values if value is not None) else SECONDS
                    width = 23 if fmt == TIME else 12
                elif name.endswith("_s"):
                    fmt, width = SECONDS, 12
                elif name == "relative_location_n":
                    fmt, width = "0", 8
                elif name == "relative_location_phase_set":
                    fmt, width = "General", 12
                elif name in {"lat", "lon"} or name.startswith(("lat_", "lon_")) or name.endswith(("_lat", "_lon", " lat", " lon")):
                    fmt, width = COORD, 12
                else:
                    continue
                for row in range(2, sheet.max_row + 1): sheet.cell(row, col).number_format = fmt
                alignment = copy(cell.alignment)
                alignment.wrap_text = True; alignment.horizontal = "center"; alignment.vertical = "center"
                cell.alignment = alignment
                sheet.column_dimensions[cell.column_letter].width = width
        aligned.column_dimensions["S"].width = 25
        aligned.column_dimensions["V"].width = 25
        wb.save(args.workbook)
    finally:
        wb.close()
    print(f"backup={backup}")

if __name__ == "__main__": main()
