#!/usr/bin/env python3
"""Extend station-acceptance conditional formats through all pair columns."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


SHEET_NAME = "Station status"


def locate_columns(sheet: object) -> tuple[int, int, int, int]:
    """Return header row, first pair column, final pair column, and final data row."""
    for row in sheet.iter_rows():
        values = [cell.value for cell in row]
        if "Station" in values and "Metric" in values:
            header_row = row[0].row
            metric_column = values.index("Metric") + 1
            pair_columns = [
                cell.column
                for cell in sheet[header_row]
                if cell.column > metric_column and cell.value not in (None, "")
            ]
            if not pair_columns:
                raise ValueError("No pair columns found after the Metric column.")
            return header_row, min(pair_columns), max(pair_columns), sheet.max_row
    raise ValueError("Could not find the Station/Metric header row.")


def extend_formats(sheet: object) -> list[tuple[str, str]]:
    header_row, first_pair_column, final_pair_column, final_row = locate_columns(sheet)
    first_data_row = header_row + 1
    records = [
        (str(conditional_range.sqref), list(sheet.conditional_formatting[conditional_range]))
        for conditional_range in sheet.conditional_formatting
    ]
    replacements: list[tuple[str, str]] = []

    sheet.conditional_formatting._cf_rules.clear()
    for range_text, rules in records:
        updated_parts = []
        for part in range_text.split():
            min_col, min_row, max_col, max_row = range_boundaries(part)
            updated_part = part
            if (
                min_col == first_pair_column
                and min_row == first_data_row
                and max_col < final_pair_column
                and max_row == final_row
            ):
                updated_part = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(final_pair_column)}{max_row}"
                )
                replacements.append((part, updated_part))
            updated_parts.append(updated_part)
        updated_range = " ".join(updated_parts)
        for rule in rules:
            sheet.conditional_formatting.add(updated_range, copy(rule))
    return replacements


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook)
    sheet = workbook[SHEET_NAME]
    replacements = extend_formats(sheet)
    if not replacements:
        raise SystemExit("No station-acceptance conditional-format range required extension.")
    workbook.save(args.workbook)
    for old, new in replacements:
        print(f"extended {old} -> {new}")


if __name__ == "__main__":
    main()
