import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

import update_station_acceptance_matrix as matrix


class StationAcceptanceMatrixTests(unittest.TestCase):
    def test_refresh_updates_results_clears_stale_cells_and_preserves_manual_shifts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "station_acceptance_matrix.xlsx"
            measurements_path = root / "run" / "phase_measurements.csv"
            output_copy = root / "run" / "station_acceptance_matrix.xlsx"
            measurements_path.parent.mkdir()

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Station status"
            sheet.cell(3, 1).value = "Source run: old"
            sheet.append([])
            sheet.append([])
            sheet.append(["Station", "Metric", "P30", "P31"])
            sheet.append(["II.HOPE P", "Status", "R", "A"])
            sheet.append(["II.HOPE P", "Differential time shift (s)", 0.5, 0.1])
            sheet.append(["II.HOPE P", "Correlation", 0.2, 0.9])
            sheet.append(["II.HOPE P", "Minimum SNR", 4.0, 5.0])
            sheet.append(["II.HOPE P", matrix.MANUAL_METRIC, 8.0, 0.0])
            sheet.append(["II.BAD P", "Status", None, None])
            sheet.append(["II.BAD P", "Differential time shift (s)", None, None])
            sheet.append(["II.BAD P", "Correlation", None, None])
            sheet.append(["II.BAD P", "Minimum SNR", None, None])
            sheet.append(["II.BAD P", matrix.MANUAL_METRIC, 0.0, None])
            sheet.append(["II.MISSING P", "Status", "X", "B"])
            sheet.append(["II.MISSING P", "Differential time shift (s)", None, None])
            sheet.append(["II.MISSING P", "Correlation", None, None])
            sheet.append(["II.MISSING P", "Minimum SNR", None, None])
            sheet.append(["II.MISSING P", matrix.MANUAL_METRIC, None, None])
            yellow_fill = PatternFill(
                fill_type="solid", start_color="FFF2CC", end_color="FFF2CC"
            )
            for cell in sheet[11]:
                cell.fill = yellow_fill
            for cell in sheet[16]:
                cell.fill = yellow_fill
            red_fill = PatternFill(
                fill_type="solid", start_color="FFC7CE", end_color="FFC7CE"
            )
            red_font = Font(color="9C0006")
            sheet.conditional_formatting.add(
                "C7:D16",
                FormulaRule(
                    formula=['AND($B7="Correlation",ISNUMBER(C7),C7<0.85)'],
                    fill=red_fill,
                    font=red_font,
                ),
            )
            workbook.save(workbook_path)
            workbook.close()

            headers = [
                "pair_label",
                "station_id",
                "phase",
                "manual_excluded",
                "good",
                "pair_residual_seconds",
                "cc",
                "automatic_pick1_snr",
                "automatic_pick2_snr",
            ]
            with measurements_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=headers)
                writer.writeheader()
                writer.writerow(
                    {
                        "pair_label": "P30",
                        "station_id": "II.HOPE",
                        "phase": "P",
                        "manual_excluded": "False",
                        "good": "True",
                        "pair_residual_seconds": "0.002",
                        "cc": "0.95",
                        "automatic_pick1_snr": "20",
                        "automatic_pick2_snr": "12",
                    }
                )
                writer.writerow(
                    {
                        "pair_label": "P31",
                        "station_id": "II.HOPE",
                        "phase": "PcP",
                        "manual_excluded": "False",
                        "good": "True",
                        "pair_residual_seconds": "-0.012",
                        "cc": "0.91",
                        "automatic_pick1_snr": "8",
                        "automatic_pick2_snr": "7",
                    }
                )
                writer.writerow(
                    {
                        "pair_label": "P30",
                        "station_id": "II.BAD",
                        "phase": "P",
                        "manual_excluded": "True",
                        "good": "False",
                        "pair_residual_seconds": "0.0",
                        "cc": "nan",
                        "automatic_pick1_snr": "0.4",
                        "automatic_pick2_snr": "1.2",
                    }
                )

            counts = matrix.refresh_station_acceptance_matrix(
                workbook_path, measurements_path, output_copy
            )
            self.assertEqual(counts["measurement_rows_updated"], 3)
            self.assertEqual(counts["station_phase_blocks_added"], 1)
            self.assertEqual(counts["conditional_format_ranges_extended"], 1)
            self.assertEqual(counts["manual_cells_preserved"], 8)
            self.assertEqual(counts["exclusions_preserved"], 2)
            self.assertEqual(counts["manual_statuses_preserved"], 2)
            self.assertEqual(counts["nan_format_rule_added"], 1)
            refreshed = load_workbook(workbook_path, data_only=True)
            sheet = refreshed["Station status"]
            def block_start(station, phase):
                return next(
                    row
                    for row in range(7, sheet.max_row + 1)
                    if sheet.cell(row, 1).value == station
                    and sheet.cell(row + 1, 1).value == phase
                )

            hope = block_start("II.HOPE", "P")
            bad = block_start("II.BAD", "P")
            missing = block_start("II.MISSING", "P")
            hope_pcp = block_start("II.HOPE", "PcP")
            self.assertEqual(sheet.cell(hope, 3).value, "A")
            self.assertIsNone(sheet.cell(hope, 4).value)
            self.assertAlmostEqual(sheet.cell(hope + 1, 3).value, 0.002)
            self.assertAlmostEqual(sheet.cell(hope + 2, 3).value, 0.95)
            self.assertAlmostEqual(sheet.cell(hope + 3, 3).value, 12.0)
            self.assertEqual(sheet.cell(hope + 4, 3).value, 8.0)
            self.assertEqual(sheet.cell(hope + 4, 4).value, 0.0)
            self.assertEqual(sheet.cell(bad, 3).value, "X")
            self.assertEqual(sheet.cell(bad + 2, 3).value, "NaN")
            self.assertEqual(sheet.cell(missing, 3).value, "X")
            self.assertEqual(sheet.cell(missing, 4).value, "B")
            self.assertEqual(sheet["A3"].value, "Source run: run")
            self.assertIn("NaN means correlation was not computed", sheet["A4"].value)
            self.assertIn("correlation at least 0.85", sheet["A4"].value)
            self.assertIn("Accepted A status cells are blue", sheet["A4"].value)
            self.assertIn("B status cells are purple", sheet["A4"].value)
            self.assertIn("non-zero manual pick-alignment shift cells are orange", sheet["A4"].value)
            correlation_rules = [
                rule
                for conditional_range in sheet.conditional_formatting
                for rule in sheet.conditional_formatting[conditional_range]
                if any('"Correlation"' in formula for formula in (rule.formula or []))
            ]
            self.assertEqual(len(correlation_rules), 3)
            self.assertTrue(
                any(
                    "C7:D26" in str(conditional_range.sqref)
                    for conditional_range in sheet.conditional_formatting
                )
            )
            self.assertTrue(
                any('="NaN"' in formula for rule in correlation_rules for formula in rule.formula)
            )
            low_rule = next(
                rule
                for rule in correlation_rules
                if any("ISNUMBER" in formula and "<0.85" in formula for formula in rule.formula)
            )
            nan_rule = next(
                rule
                for rule in correlation_rules
                if any('="NaN"' in formula for formula in rule.formula)
            )
            self.assertEqual(low_rule.dxf, nan_rule.dxf)
            exclusion_rules = [
                rule
                for conditional_range in sheet.conditional_formatting
                for rule in sheet.conditional_formatting[conditional_range]
                if any(
                    matrix.EXCLUSION_HIGHLIGHT_FORMULA_MARKER in formula
                    for formula in (rule.formula or [])
                )
            ]
            self.assertEqual(len(exclusion_rules), 1)
            self.assertTrue(exclusion_rules[0].stopIfTrue)
            self.assertEqual(exclusion_rules[0].priority, 1)
            self.assertTrue(all(cell.fill.fill_type is None for cell in sheet[hope + 4]))
            self.assertTrue(all(cell.fill.fill_type is None for cell in sheet[bad + 4]))
            self.assertEqual(sheet.cell(hope, 1).value, "II.HOPE")
            self.assertEqual(sheet.cell(hope + 1, 1).value, "P")
            self.assertIsNone(sheet.cell(hope + 2, 1).value)
            self.assertEqual(sheet.cell(hope_pcp, 1).value, "II.HOPE")
            self.assertEqual(sheet.cell(hope_pcp + 1, 1).value, "PcP")
            self.assertIsNone(sheet.cell(hope_pcp, 3).value)
            self.assertEqual(sheet.cell(hope_pcp, 4).value, "A")
            self.assertAlmostEqual(sheet.cell(hope_pcp + 1, 4).value, -0.012)
            self.assertIsNone(sheet.cell(hope_pcp + 4, 3).value)
            self.assertEqual(sheet.cell(hope_pcp + 4, 4).value, 0.0)
            passing_rules = [
                formula
                for conditional_range in sheet.conditional_formatting
                for rule in sheet.conditional_formatting[conditional_range]
                for formula in (rule.formula or [])
                if ">=" in formula or "<=0.05" in formula
            ]
            self.assertEqual(len(passing_rules), 3)
            accepted_rules = [
                rule
                for conditional_range in sheet.conditional_formatting
                for rule in sheet.conditional_formatting[conditional_range]
                if any(
                    '="Status"' in formula and '="A"' in formula
                    for formula in (rule.formula or [])
                )
            ]
            self.assertEqual(len(accepted_rules), 1)
            self.assertIn(
                matrix.ACCEPTED_BLUE_FILL,
                {
                    accepted_rules[0].dxf.fill.fgColor.rgb,
                    accepted_rules[0].dxf.fill.bgColor.rgb,
                },
            )
            manual_rules = [
                rule
                for conditional_range in sheet.conditional_formatting
                for rule in sheet.conditional_formatting[conditional_range]
                if any(
                    f'="{matrix.MANUAL_METRIC}"' in formula
                    and "ISNUMBER" in formula
                    and "<>0" in formula
                    for formula in (rule.formula or [])
                )
            ]
            self.assertEqual(len(manual_rules), 1)
            self.assertIn(
                matrix.MANUAL_ORANGE_FILL,
                {
                    manual_rules[0].dxf.fill.fgColor.rgb,
                    manual_rules[0].dxf.fill.bgColor.rgb,
                },
            )
            refreshed.close()
            snapshot = load_workbook(output_copy, data_only=True)
            self.assertEqual(snapshot["Station status"].cell(hope + 4, 3).value, 8.0)
            snapshot.close()


if __name__ == "__main__":
    unittest.main()
