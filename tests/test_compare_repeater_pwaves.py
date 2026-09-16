import io
import math
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import datetime
from pathlib import Path

import numpy as np
from obspy import Trace, UTCDateTime
from obspy.taup import TauPyModel
from openpyxl import Workbook

import compare_repeater_pwaves as crp
import fit_relative_offsets as offsets
import make_multiphase_median_outputs as multi


BASE_CONFIG = {
    "bandpass_hz": [1.0, 3.0],
    "filter_order": 4,
    "start_taper_seconds": 1.0,
    "filter_startup_exclusion_seconds": 20.0,
    "target_sampling_hz": 100.0,
    "correlation_window_seconds": [-1.0, 9.0],
    "residual_window_seconds": [-1.0, 29.0],
    "plot_window_seconds": [-10.0, 40.0],
    "lag_search_seconds": 2.0,
    "noise_chunk_seconds": 30.0,
    "noise_pre_p_guard_seconds": 10.0,
    "minimum_noise_chunks": 5,
    "clipping_repeat_count": 10,
    "large_noise_chunk_mad_multiplier": 6.0,
}


def processed(data, fs=100.0, start=0.0):
    return crp.ProcessedTrace(
        data=np.asarray(data, dtype=float),
        start_epoch=float(start),
        sampling_hz=float(fs),
        native_sampling_hz=float(fs),
        path=Path("synthetic.sac"),
        network="XX",
        station="TEST",
        location="",
        station_latitude=10.0,
        station_longitude=20.0,
        header_p_seconds_from_origin=None,
    )


class CatalogTests(unittest.TestCase):
    def make_workbook(self, path, second_lat=-56.0):
        workbook = Workbook()
        pairs = workbook.active
        pairs.title = "pairs"
        pairs.append(["label", "index1", "index2", "lat", "lon", "depth"])
        pairs.append(["P31", 723, 757, -56.0, -26.0, 70.0])
        events = workbook.create_sheet("events")
        events.append(
            [
                "INDEX",
                "TIME",
                "lat_best",
                "lon_best",
                "depth_best",
                "LAT",
                "LON",
                "DEP",
            ]
        )
        events.append(
            [723, "2004-03-23T06:20:00", -56.0, -26.0, 70.0, 1.0, 2.0, 3.0]
        )
        events.append(
            [
                757,
                "2017-06-20T12:54:35",
                second_lat,
                -26.0,
                70.0,
                4.0,
                5.0,
                6.0,
            ]
        )
        workbook.save(path)

    def test_resolves_workbook_pair_and_events(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "catalog.xlsx"
            self.make_workbook(path)
            pair = crp.resolve_catalog(path, ["P31"], 0.001, 0.1)["P31"]
            self.assertEqual((pair.event1.event_id, pair.event2.event_id), (723, 757))
            self.assertEqual(
                (pair.event1.latitude, pair.event1.longitude, pair.event1.depth_km),
                (-56.0, -26.0, 70.0),
            )

    def test_coordinate_inconsistency_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "catalog.xlsx"
            self.make_workbook(path, second_lat=-55.0)
            with self.assertRaises(crp.AnalysisError):
                crp.resolve_catalog(path, ["P31"], 0.001, 0.1)


class MatchingAndArrivalTests(unittest.TestCase):
    def test_exact_network_station_candidates_and_location_change(self):
        first = [
            {
                "path": Path("a"),
                "latitude": 1.0,
                "longitude": 2.0,
                "npts": 10,
                "location": "00",
            }
        ]
        second = [
            {
                "path": Path("b"),
                "latitude": 1.0,
                "longitude": 2.0,
                "npts": 10,
                "location": "10",
            }
        ]
        self.assertEqual(
            crp.choose_trace_pair(first, second, 0.001), (Path("a"), Path("b"))
        )
        second[0]["latitude"] = 2.0
        with self.assertRaises(crp.AnalysisError):
            crp.choose_trace_pair(first, second, 0.001)

    def test_direct_p_and_missing_direct_p(self):
        model = TauPyModel("ak135")
        event = crp.Event(1, UTCDateTime("2000-01-01"), 0.0, 0.0, 50.0)
        arrival, distance = crp.direct_p_arrival(model, event, 0.0, 30.0)
        self.assertGreater(arrival, float(event.origin))
        self.assertAlmostEqual(distance, 30.0, places=5)
        with self.assertRaises(crp.AnalysisError):
            crp.direct_p_arrival(model, event, 0.0, 160.0)

    def test_requested_phase_arrivals_returns_earliest_named_branches(self):
        model = TauPyModel("ak135")
        event = crp.Event(1, UTCDateTime("2000-01-01"), 0.0, 0.0, 50.0)
        arrivals, distance = crp.requested_phase_arrivals(
            model, event, 0.0, 30.0
        )
        self.assertEqual(set(arrivals), {"P", "pP", "sP", "PP"})
        self.assertAlmostEqual(distance, 30.0, places=5)
        self.assertLess(arrivals["P"], arrivals["pP"])
        self.assertLess(arrivals["pP"], arrivals["sP"])
        self.assertLess(arrivals["sP"], arrivals["PP"])

    def test_phase_plot_times_apply_event2_alignment_lag(self):
        plotted = crp.aligned_phase_plot_times(
            {"P": 100.0, "pP": 112.0},
            {"P": 200.0, "pP": 213.0, "PP": 260.0},
            100.0,
            200.0,
            0.5,
        )
        self.assertEqual(plotted["P"], (0.0, -0.5))
        self.assertEqual(plotted["pP"], (12.0, 12.5))
        self.assertEqual(plotted["PP"], (None, 59.5))

    def test_paired_epicentral_distance_averages_both_events(self):
        first = crp.Event(1, UTCDateTime("2000-01-01"), 0.0, 0.0, 50.0)
        second = crp.Event(2, UTCDateTime("2001-01-01"), 0.0, 2.0, 50.0)
        pair = crp.Pair("PTEST", first, second, 0.0, 1.0, 50.0)
        distance = crp.paired_epicentral_distance_degrees(
            pair, 0.0, 30.0, 0.0, 30.0
        )
        self.assertAlmostEqual(distance, 29.0, places=5)

    def test_window_bounds_are_enforced(self):
        trace = processed(np.zeros(1000), fs=100.0, start=0.0)
        with self.assertRaises(crp.AnalysisError):
            crp.extract_at_epochs(trace, np.array([-0.01, 0.0]))


class StationEvaluationModeTests(unittest.TestCase):
    def test_mode_defaults_to_selected_and_rejects_unknown_value(self):
        self.assertEqual(crp.station_evaluation_mode({}), crp.STATION_MODE_SELECTED)
        self.assertEqual(
            crp.station_evaluation_mode({"station_evaluation_mode": "ALL"}),
            crp.STATION_MODE_ALL,
        )
        with self.assertRaises(crp.AnalysisError):
            crp.station_evaluation_mode({"station_evaluation_mode": "everything"})

    def test_all_candidates_are_exact_sorted_station_intersection(self):
        first_event = crp.Event(1, UTCDateTime("2000-01-01"), 0.0, 0.0, 10.0)
        second_event = crp.Event(2, UTCDateTime("2001-01-01"), 0.0, 0.0, 10.0)
        pair = crp.Pair("PTEST", first_event, second_event, 0.0, 0.0, 10.0)
        candidates = crp.all_station_candidates(
            pair,
            {"IU.BBB": [], "IU.AAA": [], "GT.ONLY1": []},
            {"IU.AAA": [], "IU.BBB": [], "GT.ONLY2": []},
        )
        self.assertEqual(
            [candidate.station_id for candidate in candidates],
            ["IU.AAA", "IU.BBB"],
        )
        self.assertTrue(
            all(candidate.posted_correlation is None for candidate in candidates)
        )

    def test_threshold_report_includes_equal_and_greater_correlations(self):
        rows = [
            {"station_id": "XX.LOW", "new_correlation": 0.899},
            {"station_id": "XX.EQUAL", "new_correlation": 0.9},
            {"station_id": "XX.HIGH", "new_correlation": 0.95},
        ]
        report = crp.rows_meeting_correlation_threshold(rows, 0.9)
        self.assertEqual(
            [row["station_id"] for row in report], ["XX.EQUAL", "XX.HIGH"]
        )
        self.assertTrue(
            all(row["selection_correlation_threshold"] == 0.9 for row in report)
        )

    def test_plot_station_filter_includes_equal_and_greater_correlations(self):
        def station(station_id, correlation):
            return crp.StationAnalysis(
                {"station_id": station_id, "new_correlation": correlation},
                [],
                np.array([]),
                np.array([]),
                np.array([]),
                np.array([]),
                np.array([]),
                np.array([]),
                crp.NoiseEstimate([], np.array([]), 0, 0, 0, 0, []),
                crp.NoiseEstimate([], np.array([]), 0, 0, 0, 0, []),
            )

        stations = [
            station("XX.LOW", 0.849),
            station("XX.EQUAL", 0.85),
            station("XX.HIGH", 0.95),
        ]
        plotted = crp.stations_meeting_correlation_threshold(stations, 0.85)
        self.assertEqual(
            [item.result["station_id"] for item in plotted],
            ["XX.EQUAL", "XX.HIGH"],
        )

    def test_terminal_progress_reports_completed_and_total(self):
        output = io.StringIO()
        with redirect_stdout(output):
            crp.report_station_progress(0, 125)
            crp.report_station_progress(19, 125, "P34", "IU.SBA", "completed")
            crp.report_station_progress(20, 125, "P34", "IU.LVC", "completed")
            crp.report_station_progress(125, 125, "P39", "IU.QSPA", "completed")
        self.assertEqual(
            output.getvalue().splitlines(),
            [
                "Station evaluations completed: 20/125 | P34 IU.LVC | completed",
                "Station evaluations completed: 125/125 | P39 IU.QSPA | completed",
            ],
        )


class PreprocessingTests(unittest.TestCase):
    def write_sac(self, path, sampling_hz):
        duration = 80.0
        times = np.arange(int(duration * sampling_hz)) / sampling_hz
        data = np.sin(2 * np.pi * 2.0 * times).astype(np.float32)
        trace = Trace(data=data)
        trace.stats.network = "XX"
        trace.stats.station = "TEST"
        trace.stats.channel = "BHZ"
        trace.stats.sampling_rate = sampling_hz
        trace.stats.starttime = UTCDateTime("2000-01-01")
        trace.stats.sac = {"stla": 10.0, "stlo": 20.0, "b": 0.0, "t0": 30.0, "kt0": "P"}
        trace.write(str(path), format="SAC")

    def test_causal_filter_and_unequal_rates_resample_to_100_hz(self):
        with tempfile.TemporaryDirectory() as directory:
            first_path = Path(directory) / "first.sac"
            second_path = Path(directory) / "second.sac"
            self.write_sac(first_path, 20.0)
            self.write_sac(second_path, 40.0)
            first = crp.preprocess_trace(first_path, BASE_CONFIG)
            second = crp.preprocess_trace(second_path, BASE_CONFIG)
            self.assertEqual(first.sampling_hz, 100.0)
            self.assertEqual(second.sampling_hz, 100.0)
            self.assertAlmostEqual(len(first.data) / 5, 1600, delta=2)
            self.assertAlmostEqual(len(second.data) / 2.5, 3200, delta=2)
            self.assertIsNotNone(first.header_p_seconds_from_origin)


class AlignmentTests(unittest.TestCase):
    def test_requested_aic_snr_and_noise_window_defaults(self):
        self.assertEqual(multi.AUTOMATIC_PICK_MIN_SNR, 0.5)
        self.assertEqual(
            multi.AUTOMATIC_PICK_NOISE_WINDOW_SECONDS, (-30.0, -10.0)
        )

    def test_all_phases_use_configured_correlation_window(self):
        windows = multi.phase_correlation_windows(
            {"correlation_window_seconds": [-5.0, 10.0]}
        )
        self.assertEqual(set(windows), set(multi.PHASES))
        self.assertTrue(all(window == [-5.0, 10.0] for window in windows.values()))

    def test_all_phases_use_configured_plot_window(self):
        windows = multi.phase_plot_windows(
            {"plot_window_seconds": [-10.0, 30.0]}
        )
        self.assertEqual(set(windows), set(multi.PHASES))
        self.assertTrue(all(window == [-10.0, 30.0] for window in windows.values()))

    def test_phase_shift_summary_limits_are_configurable_and_validated(self):
        self.assertEqual(
            multi.phase_shift_summary_y_limits(
                {"phase_shift_summary_ylim_seconds": [-0.2, 0.2]}
            ),
            (-0.2, 0.2),
        )
        with self.assertRaises(crp.AnalysisError):
            multi.phase_shift_summary_y_limits(
                {"phase_shift_summary_ylim_seconds": [0.2, -0.2]}
            )

    def test_pair_exclusions_combine_global_and_pair_specific_lists(self):
        config = {
            "excluded_stations": ["IU.LSZ", "TRQA"],
            "excluded_stations_by_pair": {
                "P31": ["IU.QSPA", "OTAV"],
                "P30": ["PMSA"],
            },
        }
        self.assertEqual(
            multi.excluded_station_codes_for_pair(config, "P31"),
            {"LSZ", "TRQA", "QSPA", "OTAV"},
        )

    def test_pair_consistent_reference_ignores_rejected_cycle_and_requires_quorum(self):
        rows = [
            {"total_shift_seconds": 1.94, "good": True},
            {"total_shift_seconds": 2.58, "good": True},
            {"total_shift_seconds": 1.27, "good": False},
        ]
        self.assertAlmostEqual(
            multi.pair_consistent_reference_shift(rows, minimum_count=2), 2.26
        )
        rows[2]["good"] = True
        self.assertAlmostEqual(
            multi.pair_consistent_reference_shift(rows, minimum_count=2), 1.94
        )
        self.assertTrue(
            math.isnan(multi.pair_consistent_reference_shift(rows, minimum_count=4))
        )

    def test_picked_refinement_aligns_rejected_rows_but_preserves_aic_qc(self):
        accepted = {
            "automatic_pick1_accepted": True,
            "automatic_pick2_accepted": True,
        }
        rejected = {
            "automatic_pick1_accepted": True,
            "automatic_pick2_accepted": False,
        }
        self.assertTrue(
            multi.eligible_for_pair_consistent_refinement(accepted, "picked")
        )
        self.assertTrue(
            multi.eligible_for_pair_consistent_refinement(rejected, "picked")
        )
        self.assertTrue(
            multi.eligible_for_pair_consistent_refinement(rejected, "computed")
        )
        self.assertTrue(multi.passes_picked_mode_aic_qc(accepted))
        self.assertFalse(multi.passes_picked_mode_aic_qc(rejected))

    def test_signed_subsample_lag_and_no_polarity_reversal(self):
        fs = 100.0
        times = np.arange(0.0, 120.0, 1.0 / fs)
        base = (
            np.exp(-(((times - 50.3) / 0.22) ** 2))
            - 0.55 * np.exp(-(((times - 51.1) / 0.31) ** 2))
            + 0.25 * np.exp(-(((times - 53.0) / 0.5) ** 2))
        )
        true_lag = 0.237
        shifted = np.interp(times - true_lag, times, base, left=0.0, right=0.0)
        lag, correlation, boundary, _, _ = crp.signed_lag_correlation(
            processed(base, fs), processed(shifted, fs), 50.0, 50.0, [-1.0, 9.0], 2.0
        )
        self.assertAlmostEqual(lag, true_lag, delta=0.012)
        self.assertGreater(correlation, 0.999)
        self.assertFalse(boundary)
        inverted = processed(-shifted, fs)
        _, inverted_correlation, _, _, _ = crp.signed_lag_correlation(
            processed(base, fs), inverted, 50.0, 50.0, [-1.0, 9.0], 2.0
        )
        self.assertLess(inverted_correlation, 0.5)

    def test_symmetric_normalization_and_order_invariance(self):
        first = np.array([1.0, 2.0, 4.0, -2.0])
        second = np.array([2.0, 1.0, 3.0, -1.0])
        residual, scale1, scale2, rms = crp.symmetric_residual(first, second)
        swapped, swapped_scale2, swapped_scale1, swapped_rms = crp.symmetric_residual(
            second, first
        )
        self.assertTrue(np.allclose(residual, -swapped))
        self.assertAlmostEqual(rms, swapped_rms)
        noise1, noise2 = 0.3, 0.5
        r_sym = rms / math.sqrt((noise1 / scale1) ** 2 + (noise2 / scale2) ** 2)
        swapped_r_sym = swapped_rms / math.sqrt(
            (noise2 / swapped_scale2) ** 2 + (noise1 / swapped_scale1) ** 2
        )
        self.assertAlmostEqual(r_sym, swapped_r_sym)

    def test_exact_half_open_windows(self):
        correlation = crp.window_times([-1.0, 9.0], 100.0)
        residual = crp.window_times([-1.0, 29.0], 100.0)
        self.assertEqual(len(correlation), 1000)
        self.assertEqual(len(residual), 3000)
        self.assertEqual(correlation[0], -1.0)
        self.assertAlmostEqual(correlation[-1], 8.99)
        self.assertAlmostEqual(residual[-1], 28.99)

    def test_plot_window_clips_at_trace_end_without_shortening_measurement_window(self):
        fs = 10.0
        samples = np.sin(np.arange(200) / fs)
        first = processed(samples, fs=fs, start=0.0)
        second = processed(samples, fs=fs, start=0.0)

        plot_time, plot1, plot2, _, _, clipped = multi.extract_normalized_plot(
            first,
            second,
            arrival1=10.0,
            arrival2=10.0,
            lag_seconds=0.0,
            window=[-5.0, 15.0],
            normalization_window=[-2.0, 5.0],
        )

        self.assertTrue(clipped)
        self.assertEqual(plot_time[0], -5.0)
        self.assertAlmostEqual(plot_time[-1], 9.9)
        self.assertEqual(len(plot_time), 150)
        self.assertEqual(len(plot1), len(plot_time))
        self.assertEqual(len(plot2), len(plot_time))
        self.assertEqual(len(crp.window_times([-2.0, 5.0], fs)), 70)


class NoiseAndAssessmentTests(unittest.TestCase):
    def test_noise_chunking_and_robust_statistics(self):
        fs = 100.0
        # From startup-exclusion end at t=20 to P-10 at t=190 gives five chunks.
        times = np.arange(24000) / fs
        data = np.sin(2.0 * np.pi * times)
        trace = processed(data, fs)
        estimate = crp.chunk_noise(trace, 200.0, BASE_CONFIG)
        self.assertEqual(len(estimate.chunks), 5)
        self.assertEqual(estimate.usable_duration_seconds, 170.0)
        self.assertAlmostEqual(estimate.median_rms, math.sqrt(0.5), places=6)

    def test_last_noise_chunk_ends_ten_seconds_before_p(self):
        fs = 10.0
        times = np.arange(2100) / fs
        trace = processed(times, fs)
        estimate = crp.chunk_noise(trace, 200.0, BASE_CONFIG)
        self.assertEqual(len(estimate.chunks), 5)
        self.assertAlmostEqual(estimate.chunks[0][0], 40.0)
        self.assertAlmostEqual(estimate.chunks[-1][0], 160.0)
        self.assertAlmostEqual(estimate.chunks[-1][-1], 189.9)

    def test_empirical_null_uses_all_unshifted_combinations(self):
        chunks1 = [np.array([1.0, 0.0]), np.array([0.0, 1.0])]
        chunks2 = [np.array([0.0, 0.0]), np.array([1.0, 1.0])]
        null = crp.empirical_noise_null(chunks1, chunks2, 1.0, 1.0)
        expected = []
        for first in chunks1:
            for second in chunks2:
                expected.append(np.sqrt(np.mean((first - second) ** 2)))
        self.assertTrue(np.allclose(null, expected))

    def test_all_trace_assessment_states_and_minimum_chunks(self):
        null = np.arange(1.0, 101.0)
        same, threshold = crp.classify_trace(50.0, null, 5, 5, 5, [])
        different, _ = crp.classify_trace(100.0, null, 5, 5, 5, [])
        insufficient, _ = crp.classify_trace(1.0, null, 4, 5, 5, [])
        qc_failure, _ = crp.classify_trace(1.0, null, 5, 5, 5, ["clipping"])
        self.assertEqual(same, crp.ASSESS_SAME)
        self.assertEqual(different, crp.ASSESS_DIFFERENT)
        self.assertEqual(insufficient, crp.ASSESS_INDETERMINATE)
        self.assertEqual(qc_failure, crp.ASSESS_INDETERMINATE)
        self.assertGreater(threshold, 90.0)

    def test_pair_minimum_and_same_different_states(self):
        def station(observed, null):
            return crp.StationAnalysis(
                {"trace_assessment": crp.ASSESS_SAME, "residual_rms": observed},
                [],
                np.array([]),
                np.array([]),
                np.array([]),
                np.array([]),
                np.array([]),
                np.asarray(null),
                crp.NoiseEstimate([], np.array([]), 0, 0, 0, 0, []),
                crp.NoiseEstimate([], np.array([]), 0, 0, 0, 0, []),
            )

        indeterminate = crp.pair_bootstrap([station(1.0, [1.0])], 100, 1, 2)
        same = crp.pair_bootstrap([station(1.0, [2.0]), station(1.0, [2.0])], 100, 1, 2)
        different = crp.pair_bootstrap(
            [station(3.0, [2.0]), station(3.0, [2.0])], 100, 1, 2
        )
        self.assertEqual(indeterminate["pair_assessment"], crp.ASSESS_INDETERMINATE)
        self.assertEqual(same["pair_assessment"], crp.ASSESS_SAME)
        self.assertEqual(different["pair_assessment"], crp.ASSESS_DIFFERENT)


class OutputTests(unittest.TestCase):
    def test_workbook_mode_centers_event2_on_workbook_shift(self):
        self.assertEqual(
            multi.correlation_arrival_offsets("workbook", None, None, 1.25),
            (0.0, 1.25),
        )
        self.assertFalse(multi.aic_offsets_required_for_correlation("workbook"))

    def test_none_mode_centers_both_events_without_a_shift(self):
        self.assertEqual(
            multi.correlation_arrival_offsets("none", -3.0, 4.0, None),
            (0.0, 0.0),
        )
        self.assertFalse(multi.aic_offsets_required_for_correlation("none"))
        self.assertTrue(multi.aic_offsets_required_for_correlation("computed"))

    def test_workbook_common_shift_positions_measurement_references(self):
        arrival1, arrival2 = multi.measurement_phase_arrivals(
            100.0,
            102.0,
            common_origin_shift_seconds=-3.0,
            manual_shift_seconds=8.0,
        )
        self.assertEqual((arrival1, arrival2), (105.0, 107.0))

    def test_aic_snr_is_independent_of_timing_offset_acceptance(self):
        self.assertTrue(multi.picks_meet_minimum_snr(0.5, 1.2, 0.5))
        self.assertFalse(multi.picks_meet_minimum_snr(0.49, 1.2, 0.5))
        self.assertFalse(multi.picks_meet_minimum_snr(None, 1.2, 0.5))

    def test_read_workbook_pick_align_shifts(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "pairs.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "pairs"
            sheet.append(["label", "pick align shift"])
            sheet.append(["P30", -2.75])
            sheet.append(["P31", None])
            workbook.save(path)
            workbook.close()

            self.assertEqual(
                multi.read_workbook_pair_column(path, "pick align shift"),
                {"P30": -2.75},
            )

    def test_read_workbook_pair_column_requires_requested_header(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "pairs.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "pairs"
            sheet.append(["label", "new time shift"])
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(crp.AnalysisError, "pick align shift"):
                multi.read_workbook_pair_column(path, "pick align shift")

    def test_read_manual_pick_alignment_shifts(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "station_status.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Station status"
            sheet.append(["Station acceptance and measurements by event pair"])
            sheet.append([])
            sheet.append(["Station", "Metric", "P30", "P31"])
            sheet.append(["IU.ANMO P", "Status", "A", "R"])
            sheet.append(
                [None, "Manual pick alignment shift (s)", 0.125, None]
            )
            sheet.append(["IU.COLA PKP", "Status", "R", "A"])
            sheet.append(
                [None, "Manual pick alignment shift (s)", -0.2, 0.0]
            )
            workbook.save(path)
            workbook.close()

            self.assertEqual(
                multi.read_manual_pick_alignment_shifts(path),
                {
                    ("P30", "IU.ANMO", "P"): 0.125,
                    ("P30", "IU.COLA", "PKP"): -0.2,
                    ("P31", "IU.COLA", "PKP"): 0.0,
                },
            )

    def test_read_manual_pick_alignment_shifts_rejects_text_value(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "station_status.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Station status"
            sheet.append(["Station", "Metric", "P30"])
            sheet.append(
                ["IU.ANMO P", "Manual pick alignment shift (s)", "early"]
            )
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(crp.AnalysisError, "finite numeric"):
                multi.read_manual_pick_alignment_shifts(path)

    def test_read_workbook_manual_exclusions(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "station_status.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Station status"
            sheet.append(["Station", "Metric", "P30", "P31"])
            sheet.append(["IU.ANMO P", "Status", "X", "A"])
            sheet.append(["IU.COLA PKP", "Status", "R", "x"])
            workbook.save(path)
            workbook.close()

            self.assertEqual(
                multi.read_workbook_manual_exclusions(path),
                {("P30", "IU.ANMO", "P"), ("P31", "IU.COLA", "PKP")},
            )

    def test_none_mode_disables_pair_consistent_refinement(self):
        self.assertFalse(multi.eligible_for_pair_consistent_refinement({}, "none"))

    def test_l1_p_waveform_alignment_uses_accepted_p_measurements(self):
        rows = [
            {"phase": "P", "good": True, "total_shift_seconds": 1.8},
            {"phase": "P", "good": True, "total_shift_seconds": 2.0},
            {"phase": "P", "good": True, "total_shift_seconds": 8.0},
            {"phase": "P", "good": False, "total_shift_seconds": -9.0},
            {"phase": "PKP", "good": True, "total_shift_seconds": 4.0},
        ]
        self.assertEqual(multi.l1_p_waveform_alignment_shift(rows, 2), 2.0)
        self.assertTrue(math.isnan(multi.l1_p_waveform_alignment_shift(rows, 4)))

    def test_common_origin_shift_uses_high_snr_post_alignment_onsets(self):
        rows = [
            {
                "phase": "P",
                "automatic_pick1_accepted": True,
                "automatic_pick1_offset_s": -3.0,
                "automatic_pick1_snr": 8.0,
                "automatic_pick2_accepted": True,
                "automatic_pick2_offset_s": -1.0,
                "automatic_pick2_snr": 9.0,
            },
            {
                "phase": "P",
                "automatic_pick1_accepted": True,
                "automatic_pick1_offset_s": 7.0,
                "automatic_pick1_snr": 1.0,
                "automatic_pick2_accepted": True,
                "automatic_pick2_offset_s": 4.0,
                "automatic_pick2_snr": 8.0,
            },
        ]
        shift, count = multi.l1_common_high_snr_p_origin_shift(
            rows, waveform_alignment_shift_seconds=2.0, minimum_snr=5.0
        )
        self.assertEqual(shift, -3.0)
        self.assertEqual(count, 3)

    def test_common_origin_shift_ignores_manually_excluded_rows(self):
        rows = [
            {
                "phase": "P",
                "manual_excluded": False,
                "automatic_pick1_accepted": True,
                "automatic_pick1_offset_s": 0.1,
                "automatic_pick1_snr": 8.0,
                "automatic_pick2_accepted": False,
            },
            {
                "phase": "P",
                "manual_excluded": True,
                "automatic_pick1_accepted": True,
                "automatic_pick1_offset_s": 9.0,
                "automatic_pick1_snr": 8.0,
                "automatic_pick2_accepted": False,
            },
        ]
        shift, count = multi.l1_common_high_snr_p_origin_shift(
            rows, waveform_alignment_shift_seconds=0.0, minimum_snr=5.0
        )
        self.assertEqual(shift, 0.1)
        self.assertEqual(count, 1)

    def test_absolute_plot_pick_times_apply_alignment_then_common_shift(self):
        row = {
            "automatic_pick1_accepted": True,
            "automatic_pick1_offset_s": 0.2,
            "automatic_pick2_accepted": True,
            "automatic_pick2_offset_s": 2.1,
        }
        pick1, pick2 = multi.absolute_plot_pick_times(row, 1.9, -0.1)
        self.assertAlmostEqual(pick1, 0.3)
        self.assertAlmostEqual(pick2, 0.3)

    def test_absolute_plot_does_not_apply_workbook_common_shift_twice(self):
        row = {
            "automatic_pick1_accepted": True,
            "automatic_pick1_offset_s": 0.2,
            "automatic_pick2_accepted": True,
            "automatic_pick2_offset_s": 2.1,
            "aic_reference_includes_common_origin_shift": True,
        }
        pick1, pick2 = multi.absolute_plot_pick_times(row, 1.9, -3.0)
        self.assertAlmostEqual(pick1, 0.2)
        self.assertAlmostEqual(pick2, 0.2)

    def test_manual_station_phase_shift_moves_both_measurement_references(self):
        arrival1, arrival2 = multi.manually_shifted_phase_arrivals(
            100.0, 102.0, 8.0
        )
        self.assertEqual((arrival1, arrival2), (108.0, 110.0))

    def test_pkikp_below_110_degrees_is_ignored(self):
        self.assertFalse(multi.phase_is_usable_for_shift("PKiKP", 109.999))
        self.assertTrue(multi.phase_is_usable_for_shift("PKiKP", 110.0))
        self.assertTrue(multi.phase_is_usable_for_shift("P", 30.0))

    def test_pdiff_is_excluded_from_active_phases_and_preferred_fit(self):
        self.assertFalse(multi.phase_is_usable_for_shift("Pdiff", 99.999))
        self.assertFalse(multi.phase_is_usable_for_shift("Pdiff", 100.0))
        self.assertFalse(multi.phase_is_usable_for_shift("Pdiff", 109.999))
        self.assertFalse(multi.phase_is_usable_for_shift("Pdiff", 110.0))
        self.assertEqual(offsets.NO_PKIKP_PHASES, {"P", "PKP"})
        self.assertEqual(offsets.ALL_ACTIVE_PHASES, {"P", "PKP", "PKiKP"})

    def test_phase_plot_sort_key_orders_by_increasing_distance(self):
        rows = [
            {"station_id": "ZZ.Z", "epicentral_distance_degrees": 80.0},
            {"station_id": "BB.B", "epicentral_distance_degrees": 30.0},
            {"station_id": "AA.A", "epicentral_distance_degrees": 30.0},
        ]
        ordered = sorted(rows, key=multi.phase_plot_sort_key)
        self.assertEqual([row["station_id"] for row in ordered], ["AA.A", "BB.B", "ZZ.Z"])

    def test_waveform_plot_uses_compact_acceptance_labels(self):
        self.assertEqual(multi.acceptance_label({"good": True}), "Acc")
        self.assertEqual(multi.acceptance_label({"good": False}), "Rej")
        self.assertEqual(
            multi.acceptance_label({"good": False, "manual_excluded": True}), "X"
        )

    def test_excluded_event1_trace_is_black(self):
        self.assertEqual(multi.event1_trace_color({"manual_excluded": True}), "black")
        self.assertEqual(
            multi.event1_trace_color({"manual_excluded": False}), "tab:blue"
        )

    def test_shift_summary_skips_rows_without_pair_residuals(self):
        with tempfile.TemporaryDirectory() as temporary:
            multi.plot_shift_summary(
                Path(temporary),
                "P145",
                [{"phase": "P", "good": True}],
                math.nan,
                (-0.2, 0.2),
            )
            self.assertFalse(
                (Path(temporary) / "phase_plots" / "P145_phase_shift_summary.png").exists()
            )

    def test_phase_trace_information_includes_geometry_to_one_decimal(self):
        row = {
            "station_id": "GT.VNDA",
            "epicentral_distance_degrees": 58.123,
            "azimuth_degrees": 182.612,
            "cc": 0.72,
            "display_alignment_cc": 0.9829,
            "display_alignment_shift_seconds": 2.5827,
            "good": False,
        }
        self.assertEqual(
            multi.phase_trace_information(row),
            "GT.VNDA  dist=58.1\N{DEGREE SIGN}  az=182.6\N{DEGREE SIGN}  "
            "CC=0.98  resid=+2.58s  Rej",
        )

    def test_residual_format_uses_three_significant_digits(self):
        self.assertEqual(multi.format_residual_seconds(0.1), "+0.100")
        self.assertEqual(multi.format_residual_seconds(-0.01234), "-0.0123")

    def test_display_only_alignment_uses_its_best_cycle_shift(self):
        row = {
            "display_alignment_shift_seconds": 2.58,
            "fine_search_center_seconds": 1.94,
            "residual_lag_seconds": -0.13,
        }
        self.assertEqual(multi.display_shift(row), 2.58)

    def test_station_trace_label_uses_two_lines_and_compact_shift(self):
        label = crp.station_trace_label(
            {
                "station_id": "IU.QSPA",
                "epicentral_distance_degrees": 84.321,
                "azimuth_degrees": 142.678,
                "new_correlation": 0.93456,
                "lag_seconds_y_t_plus_lag": -0.1,
                "normalized_noise1": 0.1,
                "trace_assessment": crp.ASSESS_SAME,
            }
        )
        self.assertEqual(
            label,
            "IU.QSPA  dist=84.3\N{DEGREE SIGN}  az=142.7\N{DEGREE SIGN}\n"
            "CC=0.93  shift=-0.10 s",
        )
        self.assertEqual(len(label.splitlines()), 2)
        self.assertNotIn("noise", label.lower())
        self.assertNotIn("assessment", label.lower())

    def test_phase_markers_label_only_first_event(self):
        class Axis:
            def __init__(self):
                self.labels = []
                self.lines = []

            def axvline(self, arrival_time, **_kwargs):
                self.lines.append(arrival_time)

            def text(self, _x, _y, label, **_kwargs):
                self.labels.append(label)

            def get_xaxis_transform(self):
                return "xaxis"

        station = crp.StationAnalysis(
            result={},
            noise_rows=[],
            relative_plot_time=np.arange(-10.0, 80.0, 1.0),
            aligned_plot1=np.array([]),
            aligned_plot2=np.array([]),
            residual_time=np.array([]),
            residual=np.array([]),
            null_distribution=np.array([]),
            noise1=None,
            noise2=None,
            phase_plot_times={"P": (0.0, -0.5), "PcP": (25.0, 24.5)},
        )
        axis = Axis()
        crp.mark_phase_arrivals_on_overlay(axis, station)
        self.assertEqual(axis.labels, ["P"])
        self.assertEqual(axis.lines, [0.0])

    def test_plot_time_axis_limits_restore_configured_half_open_end(self):
        relative = crp.window_times([-10.0, 80.0], 100.0)
        self.assertEqual(crp.plot_time_axis_limits(relative), (-10.0, 80.0))

    def test_station_pair_trace_information_identifies_pair_and_events(self):
        row = {
            "pair_label": "P30",
            "event1": 723,
            "event2": 739,
            "epicentral_distance_degrees": 150.788,
            "cc": 0.9872,
            "pair_residual_seconds": 0.004419,
            "good": True,
        }
        label = multi.station_pair_trace_information(row)
        self.assertIn("P30", label)
        self.assertIn("723–739", label)
        self.assertIn("dist=150.8°", label)
        self.assertIn("Acc", label)

    def test_strict_json_accepts_unknown_distance_as_null(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manifest.json"
            crp.write_strict_json(
                path,
                {"exclusions": [{"epicentral_distance_degrees": None}]},
            )
            self.assertEqual(
                path.read_text(encoding="utf-8"),
                '{\n  "exclusions": [\n    {\n      "epicentral_distance_degrees": null\n    }\n  ]\n}\n',
            )

    def test_strict_json_rejects_nan_before_creating_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manifest.json"
            with self.assertRaises(ValueError):
                crp.write_strict_json(path, {"invalid": math.nan})
            self.assertFalse(path.exists())

    def test_runtime_report_labels_elapsed_and_cpu_seconds(self):
        output = io.StringIO()
        with redirect_stdout(output):
            crp.report_runtime(12.3456, 7.8912)
        self.assertEqual(
            output.getvalue().splitlines(),
            ["Elapsed time: 12.346 seconds", "CPU time: 7.891 seconds"],
        )

    def test_timestamped_output_never_overwrites(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            now = datetime(2026, 8, 7, 12, 0, 0)
            output = crp.create_timestamped_output(root, now)
            self.assertTrue((output / "pair_plots").is_dir())
            with self.assertRaises(FileExistsError):
                crp.create_timestamped_output(root, now)


if __name__ == "__main__":
    unittest.main()
