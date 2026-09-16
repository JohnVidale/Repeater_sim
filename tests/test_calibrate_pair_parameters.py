import csv
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import calibrate_pair_parameters as calibration


def make_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "pairs"
    sheet.append(
        [
            "label",
            "new_lat",
            "new_lon",
            "new_depth",
            "new time shift",
            "pick align shift",
        ]
    )
    sheet.append(["P1", 1.0, 2.0, 3.0, 4.0, 5.0])
    sheet.append(["P2", 6.0, 7.0, 8.0, 9.0, 10.0])
    workbook.save(path)
    workbook.close()


class PairCalibrationTests(unittest.TestCase):
    def test_candidate_copy_clears_failed_stale_location(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            candidate = root / "candidate.xlsx"
            make_workbook(source)
            calibration.prepare_candidate_workbook(
                source,
                candidate,
                ["P1", "P2"],
                {
                    "P1": {
                        "status": "ok",
                        "new_lat": "11",
                        "new_lon": "12",
                        "new_depth": "13",
                    },
                    "P2": {"status": "failed_insufficient_picks"},
                },
            )
            workbook = load_workbook(candidate, data_only=True)
            sheet = workbook["pairs"]
            self.assertEqual(
                [sheet.cell(2, column).value for column in range(2, 5)],
                [11.0, 12.0, 13.0],
            )
            self.assertEqual(
                [sheet.cell(3, column).value for column in range(2, 5)],
                [None, None, None],
            )
            workbook.close()

    def test_candidate_rows_require_all_location_and_timing_values(self):
        current = {
            "P1": {field: 0.0 for field in calibration.ADOPTED_FIELDS},
            "P2": {field: 0.0 for field in calibration.ADOPTED_FIELDS},
        }
        centroid_rows = {
            "P1": {
                "status": "ok",
                "new_lat": "1",
                "new_lon": "2",
                "new_depth": "3",
            },
            "P2": {"status": "failed_insufficient_picks"},
        }
        timing_rows = {
            "P1": {
                "computed_median_shift_seconds": "4",
                "common_time_shift_seconds": "5",
            },
            "P2": {
                "computed_median_shift_seconds": "6",
                "common_time_shift_seconds": "7",
            },
        }
        rows = calibration.build_candidate_rows(
            ["P1", "P2"], current, centroid_rows, timing_rows
        )
        self.assertTrue(rows[0]["ready_to_adopt"])
        self.assertFalse(rows[1]["ready_to_adopt"])
        self.assertEqual(rows[0]["change_new_time_shift"], 4.0)
        self.assertEqual(rows[0]["change_pick_align_shift"], 5.0)

    def test_adoption_writes_all_fields_and_checks_workbook_hash(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "ICevents_full.xlsx"
            make_workbook(workbook_path)
            candidate_csv = root / calibration.CANDIDATE_FILENAME
            fields = ["pair_label", "ready_to_adopt"] + [
                f"candidate_{field.replace(' ', '_')}"
                for field in calibration.ADOPTED_FIELDS
            ]
            with candidate_csv.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                writer.writeheader()
                writer.writerow(
                    {
                        "pair_label": "P1",
                        "ready_to_adopt": True,
                        "candidate_new_lat": 11,
                        "candidate_new_lon": 12,
                        "candidate_new_depth": 13,
                        "candidate_new_time_shift": 14,
                        "candidate_pick_align_shift": 15,
                    }
                )
            original_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
            (root / calibration.MANIFEST_FILENAME).write_text(
                json.dumps(
                    {
                        "canonical_workbook": str(workbook_path),
                        "canonical_workbook_sha256": original_hash,
                        "configured_pairs": ["P1"],
                    }
                )
            )
            calibration.adopt(candidate_csv)
            workbook = load_workbook(workbook_path, data_only=True)
            self.assertEqual(
                [workbook["pairs"].cell(2, column).value for column in range(2, 7)],
                [11.0, 12.0, 13.0, 14.0, 15.0],
            )
            workbook.close()
            with self.assertRaisesRegex(calibration.CalibrationError, "changed since preview"):
                calibration.adopt(candidate_csv)


if __name__ == "__main__":
    unittest.main()
