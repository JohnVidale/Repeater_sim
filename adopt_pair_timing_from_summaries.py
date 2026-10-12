#!/usr/bin/env python3
"""Adopt reviewed pair-wide timing values from median-summary CSV files."""
from __future__ import annotations
import argparse, csv
from pathlib import Path
from openpyxl import load_workbook
from aligned_pairs import aligned_rows, header_map, refresh_reference_times

def main():
    p=argparse.ArgumentParser(); p.add_argument('workbook',type=Path); p.add_argument('summaries',nargs='+',type=Path); a=p.parse_args()
    values={}
    for path in a.summaries:
        for row in csv.DictReader(path.open()):
            values[row['pair_label']]=(float(row['common_time_shift_seconds']),float(row['computed_median_shift_seconds']))
    w=load_workbook(a.workbook); s=w['aligned_pairs']; h=header_map(s)
    for label, r in aligned_rows(s).items():
        if label in values:
            pick, shift=values[label]; s.cell(r,h['pick_align_shift_s']).value=pick; s.cell(r,h['same_location_event2_minus_event1_shift_s']).value=shift; refresh_reference_times(w,r); print(label,pick,shift)
    temp=a.workbook.with_name('.'+a.workbook.name+'.tmp'); w.save(temp); w.close(); temp.replace(a.workbook)
if __name__=='__main__': main()
