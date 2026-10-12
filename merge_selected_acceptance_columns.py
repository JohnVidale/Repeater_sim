#!/usr/bin/env python3
"""Merge selected pair columns from a refreshed acceptance workbook."""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from openpyxl import load_workbook

SHEET = "Station status"


def columns(sheet, pairs):
    for row in sheet.iter_rows(max_row=10):
        values = {str(c.value).strip(): c.column for c in row if c.value is not None}
        if "Station" in values and "Metric" in values:
            missing = [p for p in pairs if p not in values]
            if missing:
                raise ValueError(f"Missing pair columns: {missing}")
            return values, row[0].row
    raise ValueError("Station/Metric header was not found")


def metric_rows(sheet, header_row):
    rows = {}
    station_phase = None
    for row in range(header_row + 1, sheet.max_row + 1):
        metric = sheet.cell(row, 2).value
        if metric == "Status":
            station = str(sheet.cell(row, 1).value or "").strip()
            phase = str(sheet.cell(row + 1, 1).value or "").strip()
            station_phase = station if " " in station else f"{station} {phase}"
        if station_phase and metric:
            rows[(station_phase, str(metric).strip())] = row
    return rows


def main():
    p = argparse.ArgumentParser()
    p.add_argument("authoritative", type=Path)
    p.add_argument("refreshed", type=Path)
    p.add_argument("snapshot", type=Path)
    p.add_argument("pairs", nargs="+")
    a = p.parse_args()
    target = load_workbook(a.authoritative)
    source = load_workbook(a.refreshed, data_only=False)
    try:
        ts, header = columns(target[SHEET], a.pairs)
        ss, source_header = columns(source[SHEET], a.pairs)
        target_rows = metric_rows(target[SHEET], header)
        source_rows = metric_rows(source[SHEET], source_header)
        for pair in a.pairs:
            for key, target_row in target_rows.items():
                source_row = source_rows.get(key)
                if source_row is not None:
                    target[SHEET].cell(target_row, ts[pair]).value = source[SHEET].cell(source_row, ss[pair]).value
        temp = a.authoritative.with_name(f".{a.authoritative.name}.tmp")
        target.save(temp)
        temp.replace(a.authoritative)
    finally:
        target.close(); source.close()
    a.snapshot.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(a.authoritative, a.snapshot)
    print(a.snapshot)


if __name__ == "__main__":
    main()
