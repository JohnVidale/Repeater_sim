#!/usr/bin/env python3
"""Plot the fraction of B statuses by the first-event year of each pair."""

from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
from openpyxl import load_workbook


def year_from_value(value: object) -> int:
    if isinstance(value, (date, datetime)):
        return value.year
    text = str(value).strip()
    if len(text) < 4 or not text[:4].isdigit():
        raise ValueError(f"Cannot determine year from {value!r}")
    return int(text[:4])


def pair_first_event_years(workbook_path: Path) -> dict[str, int]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["pairs"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        columns = {str(value): index for index, value in enumerate(header)}
        required = {"label", "date1"}
        missing = required - columns.keys()
        if missing:
            raise ValueError(f"pairs sheet is missing columns: {sorted(missing)}")
        return {
            str(row[columns["label"]]).strip(): year_from_value(row[columns["date1"]])
            for row in rows
            if row[columns["label"]] is not None
        }
    finally:
        workbook.close()


def status_counts_by_year(
    workbook_path: Path,
    first_event_year: dict[str, int],
) -> dict[int, Counter[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Station status"]
        rows = sheet.iter_rows(values_only=True)
        for _ in range(5):
            next(rows)
        header = next(rows)
        pair_labels = [str(value).strip() for value in header[2:] if value is not None]
        missing_pairs = sorted(set(pair_labels) - first_event_year.keys())
        if missing_pairs:
            raise ValueError(f"Pair labels missing from pairs sheet: {missing_pairs}")

        counts: dict[int, Counter[str]] = defaultdict(Counter)
        for row in rows:
            if str(row[1]).strip() != "Status":
                continue
            for pair_label, value in zip(pair_labels, row[2:]):
                if value is None or not str(value).strip():
                    continue
                counts[first_event_year[pair_label]][str(value).strip().upper()] += 1
        return dict(counts)
    finally:
        workbook.close()


def wilson_interval(successes: int, total: int, z: float = 1.959963984540054) -> tuple[float, float]:
    proportion = successes / total
    denominator = 1.0 + z * z / total
    center = (proportion + z * z / (2.0 * total)) / denominator
    half_width = (
        z
        * ((proportion * (1.0 - proportion) / total + z * z / (4.0 * total * total)) ** 0.5)
        / denominator
    )
    return center - half_width, center + half_width


def write_summary(path: Path, counts: dict[int, Counter[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "first_event_year",
                "b_count",
                "all_nonblank_status_count",
                "b_fraction",
                "wilson_95_lower",
                "wilson_95_upper",
            ]
        )
        for year in sorted(counts):
            total = sum(counts[year].values())
            b_count = counts[year]["B"]
            lower, upper = wilson_interval(b_count, total)
            writer.writerow([year, b_count, total, b_count / total, lower, upper])


def make_plot(path: Path, counts: dict[int, Counter[str]]) -> None:
    years = sorted(counts)
    totals = [sum(counts[year].values()) for year in years]
    b_counts = [counts[year]["B"] for year in years]
    fractions = [b_count / total for b_count, total in zip(b_counts, totals)]
    intervals = [wilson_interval(b_count, total) for b_count, total in zip(b_counts, totals)]
    lower_errors = [fraction - lower for fraction, (lower, _) in zip(fractions, intervals)]
    upper_errors = [upper - fraction for fraction, (_, upper) in zip(fractions, intervals)]

    figure, axis = plt.subplots(figsize=(9.2, 5.8), constrained_layout=True)
    axis.errorbar(
        years,
        fractions,
        yerr=[lower_errors, upper_errors],
        color="#7b3294",
        marker="o",
        markersize=7,
        linewidth=2,
        capsize=4,
        label="B fraction with 95% Wilson interval",
    )
    for year, fraction, b_count, total in zip(years, fractions, b_counts, totals):
        axis.annotate(
            f"{b_count}/{total}",
            (year, fraction),
            xytext=(0, 10),
            textcoords="offset points",
            ha="center",
            fontsize=9,
        )

    axis.set_title("Fraction of status cells flagged B by first-event year", pad=28)
    axis.text(
        0.5,
        1.01,
        "One observation is a nonblank station–phase × event-pair status cell; fraction = B / (A + R + X + B).",
        transform=axis.transAxes,
        ha="center",
        va="bottom",
        fontsize=9,
        color="0.3",
    )
    axis.set_xlabel("Year of first event in pair")
    axis.set_ylabel("Fraction with B status")
    axis.set_xticks(years)
    axis.set_ylim(0.0, max(0.13, max(upper for _, upper in intervals) * 1.17))
    axis.yaxis.set_major_formatter(plt.matplotlib.ticker.PercentFormatter(1.0, decimals=0))
    axis.grid(axis="y", color="0.86", linewidth=0.8)
    axis.spines[["top", "right"]].set_visible(False)
    axis.legend(loc="upper left", frameon=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(path.with_suffix(".png"), dpi=220)
    figure.savefig(path.with_suffix(".pdf"))
    plt.close(figure)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--events-workbook", type=Path, default=Path("ICevents_working.xlsx"))
    parser.add_argument(
        "--acceptance-workbook",
        type=Path,
        default=Path("station_acceptance_working.xlsx"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("latest_run/b_status_fraction_by_first_event_year"),
        help="Output basename; PNG, PDF, and CSV are written.",
    )
    args = parser.parse_args()

    years = pair_first_event_years(args.events_workbook)
    counts = status_counts_by_year(args.acceptance_workbook, years)
    if not counts:
        raise ValueError("No nonblank status cells were found")
    make_plot(args.output, counts)
    write_summary(args.output.with_suffix(".csv"), counts)


if __name__ == "__main__":
    main()
