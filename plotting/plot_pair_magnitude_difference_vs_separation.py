#!/usr/bin/env python3
"""Plot an event-pair magnitude statistic against repeater separation."""

from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path

import matplotlib.pyplot as plt
import openpyxl
from matplotlib.lines import Line2D


ROOT = Path(__file__).resolve().parents[1]

# Keep labels legible where several pairs have nearly identical magnitude
# differences and separations.
LABEL_OFFSETS = {
    "P31": (8, 10),
    "P34": (8, -17),
    "P112": (8, 10),
    "P321": (8, 10),
    "P355": (8, -17),
}


def csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def workbook_pair_magnitude_statistics(workbook_path: Path) -> dict[str, tuple[float, float]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        events_sheet = workbook["events"]
        event_headers = [cell.value for cell in next(events_sheet.iter_rows(max_row=1))]
        event_index = {str(header).strip(): index for index, header in enumerate(event_headers) if header is not None}
        events = {
            str(row[event_index["INDEX"]]): float(row[event_index["MAG"]])
            for row in events_sheet.iter_rows(min_row=2, values_only=True)
            if row[event_index["INDEX"]] is not None and row[event_index["MAG"]] is not None
        }
        pairs_sheet = workbook["pairs"]
        pair_headers = [cell.value for cell in next(pairs_sheet.iter_rows(max_row=1))]
        pair_index = {str(header).strip(): index for index, header in enumerate(pair_headers) if header is not None}
        statistics: dict[str, tuple[float, float]] = {}
        for row in pairs_sheet.iter_rows(min_row=2, values_only=True):
            label = row[pair_index["label"]]
            event1 = row[pair_index["index1"]]
            event2 = row[pair_index["index2"]]
            if label is None or event1 is None or event2 is None:
                continue
            try:
                magnitude1, magnitude2 = events[str(event1)], events[str(event2)]
                statistics[str(label)] = (
                    abs(magnitude2 - magnitude1),
                    0.5 * (magnitude1 + magnitude2),
                )
            except (KeyError, TypeError, ValueError):
                continue
        return statistics
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=ROOT / "analysis_config.json")
    parser.add_argument("--catalog", type=Path, default=ROOT / "ICevents_working.xlsx")
    parser.add_argument("--output-directory", type=Path, default=ROOT / "latest_run")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--magnitude-statistic",
        choices=("absolute_difference", "average"),
        default="absolute_difference",
    )
    parser.add_argument("--exclude-pairs", nargs="*", default=[])
    args = parser.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8"))
    magnitude_statistics = workbook_pair_magnitude_statistics(args.catalog)
    statistic_index = 0 if args.magnitude_statistic == "absolute_difference" else 1
    statistic_label = (
        "Absolute event-magnitude difference"
        if args.magnitude_statistic == "absolute_difference"
        else "Mean event magnitude"
    )
    separations = {
        row["pair"]: float(row["horizontal_km"])
        for row in csv_rows(
            args.output_directory
            / "median_relative_location_offsets_p_only_with_bootstrap_uncertainty.csv"
        )
        if row.get("horizontal_km")
    }
    purple, blue = set(config["pairs_purple"]), set(config["pairs_blue"])
    orange = set(config.get("pairs_orange", []))
    colors = {"orange": "#ffb000", "purple": "#6a3d9a", "blue": "#1f78b4", "other": "#555555"}
    excluded = set(args.exclude_pairs)

    figure, axis = plt.subplots(figsize=(9.5, 6.5), constrained_layout=True)
    plotted_groups: set[str] = set()
    missing: list[str] = []
    for pair in config["pairs"]:
        if pair in excluded:
            continue
        separation = separations.get(pair)
        statistics = magnitude_statistics.get(pair)
        magnitude_value = statistics[statistic_index] if statistics is not None else None
        if separation is None or magnitude_value is None or not math.isfinite(separation):
            missing.append(pair)
            continue
        group = "orange" if pair in orange else "purple" if pair in purple else "blue" if pair in blue else "other"
        plotted_groups.add(group)
        axis.scatter(separation, magnitude_value, s=86, color=colors[group], edgecolor="black", linewidth=0.6, zorder=3)
        axis.annotate(
            pair,
            (separation, magnitude_value),
            xytext=LABEL_OFFSETS.get(pair, (6, 7)),
            textcoords="offset points",
            fontsize=9,
        )

    axis.set_xlabel("Preferred horizontal repeater separation (km)")
    axis.set_ylabel(statistic_label)
    axis.set_title(f"{statistic_label} versus repeater separation")
    axis.grid(True, alpha=0.28, zorder=0)
    handles = [
        Line2D([], [], marker="o", linestyle="", markersize=10, markerfacecolor=colors[group], markeredgecolor="black", label=f"{group.title()} group")
        for group in ("orange", "purple", "blue")
        if group in plotted_groups
    ]
    if handles:
        axis.legend(handles=handles, loc="best")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(args.output, dpi=220)
    print(args.output)
    print("Not plotted: " + (", ".join(missing) if missing else "none"))


if __name__ == "__main__":
    main()
