#!/usr/bin/env python3
"""Refresh a station acceptance matrix from one multiphase measurement table."""

from __future__ import annotations

import argparse
import csv
import math
import os
import shutil
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SHEET_NAME = "Station status"
CALCULATED_METRICS = (
    "Status",
    "Differential time shift (s)",
    "Correlation",
    "Minimum SNR",
)
MANUAL_METRIC = "Manual pick alignment shift (s)"
UNCOMPUTED_CORRELATION = "NaN"
EXCLUSION_HIGHLIGHT_FORMULA_MARKER = "LOOKUP(2,1/("
ACCEPTED_BLUE_FILL = "FFBDD7EE"
ACCEPTED_BLUE_FONT = "FF1F4E78"
MANUAL_ORANGE_FILL = "FFF4B183"
MANUAL_ORANGE_FONT = "FF833C0C"


class MatrixUpdateError(RuntimeError):
    """Raised when the workbook cannot be refreshed without ambiguity."""


def parse_bool(value: Any) -> bool:
    return str(value).strip().lower() == "true"


def finite_float(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def minimum_snr(row: dict[str, str]) -> float | None:
    first = finite_float(row.get("automatic_pick1_snr"))
    second = finite_float(row.get("automatic_pick2_snr"))
    return min(first, second) if first is not None and second is not None else None


def correlation_value(row: dict[str, str]) -> float | str:
    value = finite_float(row.get("cc"))
    return value if value is not None else UNCOMPUTED_CORRELATION


def measurement_values(row: dict[str, str]) -> dict[str, str | float | None]:
    status = "X" if parse_bool(row.get("manual_excluded")) else (
        "A" if parse_bool(row.get("good")) else "R"
    )
    return {
        "Status": status,
        "Differential time shift (s)": finite_float(
            row.get("pair_residual_seconds")
        ),
        "Correlation": correlation_value(row),
        "Minimum SNR": minimum_snr(row),
    }


def ensure_uncomputed_correlation_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Color NaN correlations with the existing low-correlation style."""
    low_correlation_rule = None
    low_correlation_range = None
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            formulas = [str(formula) for formula in (rule.formula or [])]
            if any(
                '"Correlation"' in formula and '="NaN"' in formula
                for formula in formulas
            ):
                return False
            if any(
                '"Correlation"' in formula and "ISNUMBER" in formula
                for formula in formulas
            ):
                low_correlation_rule = rule
                low_correlation_range = conditional_range

    if (
        low_correlation_rule is None
        or low_correlation_rule.dxf is None
        or low_correlation_range is None
        or not pair_columns
    ):
        return False

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}="Correlation",'
        f'{first_pair_letter}{first_data_row}="{UNCOMPUTED_CORRELATION}")'
    )
    sheet.conditional_formatting.add(
        str(low_correlation_range.sqref),
        Rule(
            type="expression",
            dxf=copy(low_correlation_rule.dxf),
            formula=[formula],
        ),
    )
    return True


def ensure_exclusion_highlight_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Highlight the four calculated cells for each X station-pair entry."""
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if any(
                EXCLUSION_HIGHLIGHT_FORMULA_MARKER in str(formula)
                and '="X"' in str(formula)
                for formula in (rule.formula or [])
            ):
                return False
    if not pair_columns:
        return False

    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if rule.priority is not None:
                rule.priority += 1

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}<>"{MANUAL_METRIC}",'
        f'LOOKUP(2,1/(${metric_letter}${first_data_row}:${metric_letter}{first_data_row}'
        f'="Status"),{first_pair_letter}${first_data_row}:{first_pair_letter}'
        f'{first_data_row})="X")'
    )
    rule = FormulaRule(
        formula=[formula],
        fill=PatternFill(
            fill_type="solid", start_color="FFFFE699", end_color="FFFFE699"
        ),
        font=Font(color="FF000000"),
        stopIfTrue=True,
    )
    rule.priority = 1
    sheet.conditional_formatting.add(
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}",
        rule,
    )
    return True


def clear_manual_row_backgrounds(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    metric_rows: dict[tuple[str, str], int],
) -> int:
    """Remove direct background fills from every manual-shift row."""
    cleared = 0
    no_fill = PatternFill()
    for (_, metric), row_number in metric_rows.items():
        if metric != MANUAL_METRIC:
            continue
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row_number, column)
            if cell.fill.fill_type is not None:
                cleared += 1
            cell.fill = copy(no_fill)
    return cleared


def ensure_accepted_status_blue_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Color accepted Status cells blue, updating an existing A rule if present."""
    if not pair_columns:
        return False
    blue_fill = PatternFill(
        fill_type="solid",
        start_color=ACCEPTED_BLUE_FILL,
        end_color=ACCEPTED_BLUE_FILL,
    )
    blue_font = Font(color=ACCEPTED_BLUE_FONT)
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            formulas = [str(formula) for formula in (rule.formula or [])]
            if any(
                '="Status"' in formula and '="A"' in formula
                for formula in formulas
            ):
                if rule.dxf is None:
                    rule.dxf = openpyxl.styles.differential.DifferentialStyle()
                rule.dxf.fill = copy(blue_fill)
                rule.dxf.font = copy(blue_font)
                return False

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}="Status",'
        f'{first_pair_letter}{first_data_row}="A")'
    )
    sheet.conditional_formatting.add(
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}",
        FormulaRule(formula=[formula], fill=blue_fill, font=blue_font),
    )
    return True


def ensure_nonzero_manual_shift_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Color non-zero numeric manual pick-alignment shifts orange."""
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if any(
                f'="{MANUAL_METRIC}"' in str(formula)
                and "ISNUMBER" in str(formula)
                and "<>0" in str(formula)
                for formula in (rule.formula or [])
            ):
                return False
    if not pair_columns:
        return False

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}="{MANUAL_METRIC}",'
        f'ISNUMBER({first_pair_letter}{first_data_row}),'
        f'{first_pair_letter}{first_data_row}<>0)'
    )
    sheet.conditional_formatting.add(
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}",
        FormulaRule(
            formula=[formula],
            fill=PatternFill(
                fill_type="solid",
                start_color=MANUAL_ORANGE_FILL,
                end_color=MANUAL_ORANGE_FILL,
            ),
            font=Font(color=MANUAL_ORANGE_FONT),
        ),
    )
    return True


def ensure_passing_threshold_formats(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> int:
    """Color numeric measurements green when they meet display thresholds."""
    existing_formulas = {
        str(formula)
        for conditional_range in sheet.conditional_formatting
        for rule in sheet.conditional_formatting[conditional_range]
        for formula in (rule.formula or [])
    }
    if not pair_columns:
        return 0
    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    target_range = (
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}"
    )
    conditions = (
        ("Differential time shift (s)", f"ABS({first_pair_letter}{first_data_row})<=0.05"),
        ("Correlation", f"{first_pair_letter}{first_data_row}>=0.85"),
        ("Minimum SNR", f"{first_pair_letter}{first_data_row}>=1"),
    )
    added = 0
    for metric, threshold_test in conditions:
        formula = (
            f'AND(${metric_letter}{first_data_row}="{metric}",'
            f"ISNUMBER({first_pair_letter}{first_data_row}),{threshold_test})"
        )
        if formula in existing_formulas:
            continue
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[formula],
                fill=PatternFill(
                    fill_type="solid", start_color="FFC6EFCE", end_color="FFC6EFCE"
                ),
                font=Font(color="FF006100"),
            ),
        )
        added += 1
    return added


def show_station_only_on_status_rows(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    station_column: int,
    metric_rows: dict[tuple[str, str], int],
) -> int:
    """Keep station-phase text on Status rows and blank it on detail rows."""
    changed = 0
    for (station_phase, metric), row_number in metric_rows.items():
        cell = sheet.cell(row_number, station_column)
        expected = station_phase if metric == "Status" else None
        if cell.value != expected:
            changed += 1
        cell.value = expected
    return changed


def read_measurements(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    measurements: dict[tuple[str, str], dict[str, str]] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            pair = str(row.get("pair_label", "")).strip()
            station = str(row.get("station_id", "")).strip()
            phase = str(row.get("phase", "")).strip()
            if not pair or not station or not phase:
                continue
            key = (f"{station} {phase}", pair)
            if key in measurements:
                raise MatrixUpdateError(
                    f"Duplicate phase measurement for {station} {phase} {pair}"
                )
            measurements[key] = row
    return measurements


def find_header_row(sheet: openpyxl.worksheet.worksheet.Worksheet) -> int:
    for row in sheet.iter_rows():
        values = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in row]
        if "station" in values and "metric" in values:
            return row[0].row
    raise MatrixUpdateError("Station acceptance matrix lacks Station and Metric headers")


def refresh_station_acceptance_matrix(
    matrix_path: Path,
    measurements_path: Path,
    output_copy: Path,
) -> dict[str, int]:
    measurements = read_measurements(measurements_path)
    workbook = openpyxl.load_workbook(matrix_path)
    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise MatrixUpdateError(f"Workbook lacks required {SHEET_NAME!r} sheet")
    sheet = workbook[SHEET_NAME]
    header_row = find_header_row(sheet)
    headers = {
        str(cell.value).strip(): cell.column
        for cell in sheet[header_row]
        if cell.value is not None and str(cell.value).strip()
    }
    if "Station" not in headers or "Metric" not in headers:
        workbook.close()
        raise MatrixUpdateError("Station acceptance matrix headers are malformed")
    station_column = headers["Station"]
    metric_column = headers["Metric"]
    pair_columns = {
        label: column
        for label, column in headers.items()
        if label not in {"Station", "Metric"}
    }

    metric_rows: dict[tuple[str, str], int] = {}
    manual_values_before: dict[tuple[str, str], Any] = {}
    exclusions_before: set[tuple[str, str]] = set()
    current_station_phase: str | None = None
    for row_number in range(header_row + 1, sheet.max_row + 1):
        station_phase_value = sheet.cell(row_number, station_column).value
        metric = sheet.cell(row_number, metric_column).value
        if station_phase_value is not None and str(station_phase_value).strip():
            current_station_phase = str(station_phase_value).strip()
        if current_station_phase is None or metric is None:
            continue
        key = (current_station_phase, str(metric).strip())
        if key in metric_rows:
            workbook.close()
            raise MatrixUpdateError(f"Duplicate matrix row for {key[0]} / {key[1]}")
        metric_rows[key] = row_number
        if key[1] == MANUAL_METRIC:
            for pair, column in pair_columns.items():
                manual_values_before[(key[0], pair)] = sheet.cell(
                    row_number, column
                ).value
        elif key[1] == "Status":
            for pair, column in pair_columns.items():
                if str(sheet.cell(row_number, column).value).strip().upper() == "X":
                    exclusions_before.add((key[0], pair))

    for (station_phase, metric), row_number in metric_rows.items():
        if metric not in CALCULATED_METRICS:
            continue
        for column in pair_columns.values():
            sheet.cell(row_number, column).value = None

    updated_measurements = 0
    for (station_phase, pair), measurement in measurements.items():
        if pair not in pair_columns:
            workbook.close()
            raise MatrixUpdateError(f"Measurement pair {pair} has no matrix column")
        values = measurement_values(measurement)
        for metric, value in values.items():
            row_number = metric_rows.get((station_phase, metric))
            if row_number is None:
                workbook.close()
                raise MatrixUpdateError(
                    f"Measurement {station_phase} has no {metric!r} matrix row"
                )
            sheet.cell(row_number, pair_columns[pair]).value = value
        updated_measurements += 1

    for station_phase, pair in exclusions_before:
        status_row = metric_rows.get((station_phase, "Status"))
        if status_row is None or pair not in pair_columns:
            workbook.close()
            raise MatrixUpdateError(
                f"Stored exclusion {station_phase} {pair} has no Status cell"
            )
        sheet.cell(status_row, pair_columns[pair]).value = "X"

    nan_format_rule_added = ensure_uncomputed_correlation_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    exclusion_format_rule_added = ensure_exclusion_highlight_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    manual_background_cells_cleared = clear_manual_row_backgrounds(
        sheet,
        metric_rows,
    )
    accepted_blue_rule_added = ensure_accepted_status_blue_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    manual_orange_rule_added = ensure_nonzero_manual_shift_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    passing_format_rules_added = ensure_passing_threshold_formats(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    station_labels_changed = show_station_only_on_status_rows(
        sheet,
        station_column,
        metric_rows,
    )

    source_label = f"Source run: {measurements_path.parent.name}"
    sheet.cell(3, 1).value = source_label
    sheet.cell(4, 1).value = (
        "All station-phase rows from the 16-pair run are retained in distance "
        "order. Enter a manual pick-alignment shift in its labeled row; 0.000 "
        "means no manual adjustment. An X is measured and plotted but excluded "
        "from differential-location estimation; its calculated cells are yellow. "
        "Accepted A status cells are blue, and non-zero manual pick-alignment "
        "shift cells are orange. "
        "Green numeric cells meet the display thresholds: absolute shift at most "
        "0.05 s, correlation at least 0.85, and minimum SNR at least 1.0. Blank "
        "cells have no measurement; NaN means correlation was not computed."
    )
    for (station_phase, pair), prior_value in manual_values_before.items():
        row_number = metric_rows[(station_phase, MANUAL_METRIC)]
        if sheet.cell(row_number, pair_columns[pair]).value != prior_value:
            workbook.close()
            raise MatrixUpdateError(
                f"Manual shift changed unexpectedly for {station_phase} {pair}"
            )

    temporary_path = matrix_path.with_name(f".{matrix_path.name}.tmp")
    try:
        workbook.save(temporary_path)
    finally:
        workbook.close()
    os.replace(temporary_path, matrix_path)
    output_copy.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(matrix_path, output_copy)
    return {
        "measurement_rows_updated": updated_measurements,
        "manual_cells_preserved": len(manual_values_before),
        "nan_format_rule_added": int(nan_format_rule_added),
        "exclusion_format_rule_added": int(exclusion_format_rule_added),
        "manual_background_cells_cleared": manual_background_cells_cleared,
        "accepted_blue_rule_added": int(accepted_blue_rule_added),
        "manual_orange_rule_added": int(manual_orange_rule_added),
        "passing_format_rules_added": passing_format_rules_added,
        "station_labels_changed": station_labels_changed,
        "exclusions_preserved": len(exclusions_before),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--measurements", type=Path, required=True)
    parser.add_argument("--output-copy", type=Path, required=True)
    args = parser.parse_args()
    counts = refresh_station_acceptance_matrix(
        args.matrix.resolve(),
        args.measurements.resolve(),
        args.output_copy.resolve(),
    )
    print(f"updated={counts['measurement_rows_updated']}")
    print(f"manual_cells_preserved={counts['manual_cells_preserved']}")
    print(args.matrix.resolve())
    print(args.output_copy.resolve())


if __name__ == "__main__":
    main()
