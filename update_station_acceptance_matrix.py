#!/usr/bin/env python3
"""Refresh a station acceptance matrix from one multiphase measurement table."""

from __future__ import annotations

import argparse
import csv
import math
import os
import shutil
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries


SHEET_NAME = "Station status"
CALCULATED_METRICS = (
    "Status",
    "Differential time shift (s)",
    "Correlation",
    "Minimum SNR",
)
MANUAL_METRIC = "Manual pick alignment shift (s)"
UNCOMPUTED_CORRELATION = "NaN"
EXCLUSION_HIGHLIGHT_FORMULA_MARKER = "LOOKUP(2,1/("
ACCEPTED_BLUE_FILL = "FFBDD7EE"
ACCEPTED_BLUE_FONT = "FF1F4E78"
MANUAL_ORANGE_FILL = "FFF4B183"
MANUAL_ORANGE_FONT = "FF833C0C"
B_STATUS_FILL = "FFE4DFEC"
B_STATUS_FONT = "FF5F497A"
PHASE_DISPLAY_ORDER = ("P", "PcP", "ScP", "PKiKP", "PKP")


class MatrixUpdateError(RuntimeError):
    """Raised when the workbook cannot be refreshed without ambiguity."""


def parse_bool(value: Any) -> bool:
    return str(value).strip().lower() == "true"


def finite_float(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def minimum_snr(row: dict[str, str]) -> float | None:
    first = finite_float(row.get("automatic_pick1_snr"))
    second = finite_float(row.get("automatic_pick2_snr"))
    return min(first, second) if first is not None and second is not None else None


def correlation_value(row: dict[str, str]) -> float | str:
    value = finite_float(row.get("cc"))
    return value if value is not None else UNCOMPUTED_CORRELATION


def measurement_values(row: dict[str, str]) -> dict[str, str | float | None]:
    manual_status = str(row.get("manual_status") or "").strip().upper()
    status = (
        manual_status
        if manual_status in {"X", "B"}
        else "X"
        if parse_bool(row.get("manual_excluded"))
        else "A"
        if parse_bool(row.get("good"))
        else "R"
    )
    return {
        "Status": status,
        "Differential time shift (s)": finite_float(
            row.get("pair_residual_seconds")
        ),
        "Correlation": correlation_value(row),
        "Minimum SNR": minimum_snr(row),
    }


def ensure_uncomputed_correlation_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Color NaN correlations with the existing low-correlation style."""
    low_correlation_rule = None
    low_correlation_range = None
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            formulas = [str(formula) for formula in (rule.formula or [])]
            if any(
                '"Correlation"' in formula and '="NaN"' in formula
                for formula in formulas
            ):
                return False
            if any(
                '"Correlation"' in formula and "ISNUMBER" in formula
                for formula in formulas
            ):
                low_correlation_rule = rule
                low_correlation_range = conditional_range

    if (
        low_correlation_rule is None
        or low_correlation_rule.dxf is None
        or low_correlation_range is None
        or not pair_columns
    ):
        return False

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}="Correlation",'
        f'{first_pair_letter}{first_data_row}="{UNCOMPUTED_CORRELATION}")'
    )
    sheet.conditional_formatting.add(
        str(low_correlation_range.sqref),
        Rule(
            type="expression",
            dxf=copy(low_correlation_rule.dxf),
            formula=[formula],
        ),
    )
    return True


def ensure_exclusion_highlight_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Highlight the four calculated cells for each X station-pair entry."""
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if any(
                EXCLUSION_HIGHLIGHT_FORMULA_MARKER in str(formula)
                and '="X"' in str(formula)
                for formula in (rule.formula or [])
            ):
                return False
    if not pair_columns:
        return False

    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if rule.priority is not None:
                rule.priority += 1

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}<>"{MANUAL_METRIC}",'
        f'LOOKUP(2,1/(${metric_letter}${first_data_row}:${metric_letter}{first_data_row}'
        f'="Status"),{first_pair_letter}${first_data_row}:{first_pair_letter}'
        f'{first_data_row})="X")'
    )
    rule = FormulaRule(
        formula=[formula],
        fill=PatternFill(
            fill_type="solid", start_color="FFFFE699", end_color="FFFFE699"
        ),
        font=Font(color="FF000000"),
        stopIfTrue=True,
    )
    rule.priority = 1
    sheet.conditional_formatting.add(
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}",
        rule,
    )
    return True


def clear_manual_row_backgrounds(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    metric_rows: dict[tuple[str, str], int],
) -> int:
    """Remove direct background fills from every manual-shift row."""
    cleared = 0
    no_fill = PatternFill()
    for (_, metric), row_number in metric_rows.items():
        if metric != MANUAL_METRIC:
            continue
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row_number, column)
            if cell.fill.fill_type is not None:
                cleared += 1
            cell.fill = copy(no_fill)
    return cleared


def ensure_accepted_status_blue_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Color accepted Status cells blue, updating an existing A rule if present."""
    if not pair_columns:
        return False
    blue_fill = PatternFill(
        fill_type="solid",
        start_color=ACCEPTED_BLUE_FILL,
        end_color=ACCEPTED_BLUE_FILL,
    )
    blue_font = Font(color=ACCEPTED_BLUE_FONT)
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            formulas = [str(formula) for formula in (rule.formula or [])]
            if any(
                '="Status"' in formula and '="A"' in formula
                for formula in formulas
            ):
                if rule.dxf is None:
                    rule.dxf = openpyxl.styles.differential.DifferentialStyle()
                rule.dxf.fill = copy(blue_fill)
                rule.dxf.font = copy(blue_font)
                return False

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}="Status",'
        f'{first_pair_letter}{first_data_row}="A")'
    )
    sheet.conditional_formatting.add(
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}",
        FormulaRule(formula=[formula], fill=blue_fill, font=blue_font),
    )
    return True


def ensure_b_status_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Color persistent B status cells purple."""
    if not pair_columns:
        return False
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if any(
                '="Status"' in str(formula) and '="B"' in str(formula)
                for formula in (rule.formula or [])
            ):
                return False
    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}="Status",'
        f'{first_pair_letter}{first_data_row}="B")'
    )
    sheet.conditional_formatting.add(
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}",
        FormulaRule(
            formula=[formula],
            fill=PatternFill(
                fill_type="solid", start_color=B_STATUS_FILL, end_color=B_STATUS_FILL
            ),
            font=Font(color=B_STATUS_FONT),
        ),
    )
    return True


def ensure_nonzero_manual_shift_format(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> bool:
    """Color non-zero numeric manual pick-alignment shifts orange."""
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if any(
                f'="{MANUAL_METRIC}"' in str(formula)
                and "ISNUMBER" in str(formula)
                and "<>0" in str(formula)
                for formula in (rule.formula or [])
            ):
                return False
    if not pair_columns:
        return False

    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    formula = (
        f'AND(${metric_letter}{first_data_row}="{MANUAL_METRIC}",'
        f'ISNUMBER({first_pair_letter}{first_data_row}),'
        f'{first_pair_letter}{first_data_row}<>0)'
    )
    sheet.conditional_formatting.add(
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}",
        FormulaRule(
            formula=[formula],
            fill=PatternFill(
                fill_type="solid",
                start_color=MANUAL_ORANGE_FILL,
                end_color=MANUAL_ORANGE_FILL,
            ),
            font=Font(color=MANUAL_ORANGE_FONT),
        ),
    )
    return True


def ensure_passing_threshold_formats(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    metric_column: int,
    pair_columns: dict[str, int],
) -> int:
    """Color numeric measurements green when they meet display thresholds."""
    existing_formulas = {
        str(formula)
        for conditional_range in sheet.conditional_formatting
        for rule in sheet.conditional_formatting[conditional_range]
        for formula in (rule.formula or [])
    }
    if not pair_columns:
        return 0
    first_data_row = header_row + 1
    metric_letter = get_column_letter(metric_column)
    first_pair_letter = get_column_letter(min(pair_columns.values()))
    last_pair_letter = get_column_letter(max(pair_columns.values()))
    target_range = (
        f"{first_pair_letter}{first_data_row}:{last_pair_letter}{sheet.max_row}"
    )
    conditions = (
        ("Differential time shift (s)", f"ABS({first_pair_letter}{first_data_row})<=0.05"),
        ("Correlation", f"{first_pair_letter}{first_data_row}>=0.85"),
        ("Minimum SNR", f"{first_pair_letter}{first_data_row}>=1"),
    )
    added = 0
    for metric, threshold_test in conditions:
        formula = (
            f'AND(${metric_letter}{first_data_row}="{metric}",'
            f"ISNUMBER({first_pair_letter}{first_data_row}),{threshold_test})"
        )
        if formula in existing_formulas:
            continue
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[formula],
                fill=PatternFill(
                    fill_type="solid", start_color="FFC6EFCE", end_color="FFC6EFCE"
                ),
                font=Font(color="FF006100"),
            ),
        )
        added += 1
    return added


def show_station_only_on_status_rows(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    station_column: int,
    metric_rows: dict[tuple[str, str], int],
) -> int:
    """Keep station-phase text on Status rows and blank it on detail rows."""
    changed = 0
    for (station_phase, metric), row_number in metric_rows.items():
        cell = sheet.cell(row_number, station_column)
        expected = station_phase if metric == "Status" else None
        if cell.value != expected:
            changed += 1
        cell.value = expected
    return changed


def read_measurements(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    measurements: dict[tuple[str, str], dict[str, str]] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            pair = str(row.get("pair_label", "")).strip()
            station = str(row.get("station_id", "")).strip()
            phase = str(row.get("phase", "")).strip()
            if not pair or not station or not phase:
                continue
            key = (f"{station} {phase}", pair)
            if key in measurements:
                raise MatrixUpdateError(
                    f"Duplicate phase measurement for {station} {phase} {pair}"
                )
            measurements[key] = row
    return measurements


def find_header_row(sheet: openpyxl.worksheet.worksheet.Worksheet) -> int:
    for row in sheet.iter_rows():
        values = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in row]
        if "station" in values and "metric" in values:
            return row[0].row
    raise MatrixUpdateError("Station acceptance matrix lacks Station and Metric headers")


def add_missing_pair_columns(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    pair_columns: dict[str, int],
    measurements: dict[tuple[str, str], dict[str, str]],
) -> list[str]:
    """Append blank, styled pair columns needed by a newly reviewed pair.

    This is intentionally additive: existing pair columns and their manual
    statuses are never reordered or overwritten.  New manual-alignment cells
    are initialized later, and only where a trace was measured.
    """
    missing_pairs = sorted({pair for _, pair in measurements} - set(pair_columns))
    if not missing_pairs:
        return []
    if not pair_columns:
        raise MatrixUpdateError("Station acceptance matrix has no pair-column template")
    template_column = max(pair_columns.values())
    source_width = sheet.column_dimensions[get_column_letter(template_column)].width
    for pair in missing_pairs:
        target_column = sheet.max_column + 1
        for row_number in range(1, sheet.max_row + 1):
            source = sheet.cell(row_number, template_column)
            target = sheet.cell(row_number, target_column)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.alignment:
                target.alignment = copy(source.alignment)
            if source.protection:
                target.protection = copy(source.protection)
            target.value = None
        if source_width is not None:
            sheet.column_dimensions[get_column_letter(target_column)].width = source_width
        sheet.cell(header_row, target_column).value = pair
        pair_columns[pair] = target_column
    return missing_pairs


def add_missing_station_phase_blocks(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    station_column: int,
    metric_column: int,
    pair_columns: dict[str, int],
    measurements: dict[tuple[str, str], dict[str, str]],
) -> int:
    """Append editable metric blocks for station-phases newly present in a run.

    The acceptance workbook is authoritative for manual statuses and pick shifts,
    but a newly examined phase has no pre-existing rows to hold those choices.
    Copying the first available row for each metric preserves the workbook's
    layout and number formats while leaving every calculated value blank.
    """
    existing = {
        station_phase_for_status_row(
            sheet, row_number, station_column, metric_column
        )
        for row_number in range(header_row + 1, sheet.max_row + 1)
        if str(sheet.cell(row_number, metric_column).value).strip() == "Status"
        and sheet.cell(row_number, station_column).value is not None
        and str(sheet.cell(row_number, station_column).value).strip()
    }
    required = sorted({station_phase for station_phase, _ in measurements})
    missing = [station_phase for station_phase in required if station_phase not in existing]
    if not missing:
        return 0

    templates: dict[str, int] = {}
    for row_number in range(header_row + 1, sheet.max_row + 1):
        metric = sheet.cell(row_number, metric_column).value
        if metric is not None and str(metric).strip() in (*CALCULATED_METRICS, MANUAL_METRIC):
            templates.setdefault(str(metric).strip(), row_number)
    missing_templates = [
        metric for metric in (*CALCULATED_METRICS, MANUAL_METRIC) if metric not in templates
    ]
    if missing_templates:
        raise MatrixUpdateError(
            "Station acceptance matrix lacks template rows for "
            + ", ".join(missing_templates)
        )

    for station_phase in missing:
        for metric in (*CALCULATED_METRICS, MANUAL_METRIC):
            source_row = templates[metric]
            target_row = sheet.max_row + 1
            for column in range(1, sheet.max_column + 1):
                source = sheet.cell(source_row, column)
                target = sheet.cell(target_row, column)
                if source.has_style:
                    target._style = copy(source._style)
                if source.number_format:
                    target.number_format = source.number_format
                if source.alignment:
                    target.alignment = copy(source.alignment)
                if source.protection:
                    target.protection = copy(source.protection)
            sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
            sheet.cell(target_row, station_column).value = (
                station_phase if metric == "Status" else None
            )
            sheet.cell(target_row, metric_column).value = metric
            for pair, column in pair_columns.items():
                sheet.cell(target_row, column).value = (
                    0.0
                    if metric == MANUAL_METRIC and (station_phase, pair) in measurements
                    else None
                )
    return len(missing)


def extend_phase_conditional_formats(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    pair_columns: dict[str, int],
) -> int:
    """Extend existing phase-table conditional formats through appended rows."""
    if not pair_columns:
        return 0
    first_data_row = header_row + 1
    first_pair_column = min(pair_columns.values())
    last_pair_column = max(pair_columns.values())
    final_row = sheet.max_row
    records = [
        (str(conditional_range.sqref), list(sheet.conditional_formatting[conditional_range]))
        for conditional_range in sheet.conditional_formatting
    ]
    updated = 0
    replacement_records: list[tuple[str, list[Rule]]] = []
    for range_text, rules in records:
        replacement_ranges: list[str] = []
        changed = False
        for part in range_text.split():
            min_col, min_row, max_col, max_row = range_boundaries(part)
            if (
                min_row == first_data_row
                and max_row != final_row
                and min_col >= first_pair_column
                and max_col <= last_pair_column
            ):
                part = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{final_row}"
                )
                changed = True
            replacement_ranges.append(part)
        replacement_records.append((" ".join(replacement_ranges), rules))
        updated += int(changed)
    if updated:
        sheet.conditional_formatting._cf_rules.clear()
        for range_text, rules in replacement_records:
            for rule in rules:
                sheet.conditional_formatting.add(range_text, rule)
    return updated


def station_phase_for_status_row(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    row_number: int,
    station_column: int,
    metric_column: int,
) -> str:
    """Read either the legacy combined label or the three-line display label."""
    station_value = str(sheet.cell(row_number, station_column).value or "").strip()
    if not station_value:
        raise MatrixUpdateError(f"Status row {row_number} has no station label")
    if " " in station_value:
        return station_value
    phase_row = row_number + 1
    phase = str(sheet.cell(phase_row, station_column).value or "").strip()
    if sheet.cell(phase_row, metric_column).value != "Differential time shift (s)" or not phase:
        raise MatrixUpdateError(
            f"Status row {row_number} lacks the phase label in the following row"
        )
    return f"{station_value} {phase}"


def remove_obsolete_phase_blocks(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    station_column: int,
    metric_column: int,
    measurements: dict[tuple[str, str], dict[str, str]],
    phases: set[str],
) -> int:
    """Remove phase blocks having no measurements in the current phase scope."""
    required = {station_phase for station_phase, _ in measurements}
    status_rows = [
        row_number
        for row_number in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(row_number, metric_column).value == "Status"
    ]
    removals: list[tuple[int, int]] = []
    for index, start in enumerate(status_rows):
        end = status_rows[index + 1] if index + 1 < len(status_rows) else sheet.max_row + 1
        station_phase = station_phase_for_status_row(
            sheet, start, station_column, metric_column
        )
        _, phase = station_phase.rsplit(" ", 1)
        if phase in phases and station_phase not in required:
            removals.append((start, end - start))
    for start, count in reversed(removals):
        sheet.delete_rows(start, count)
    return len(removals)


def phase_distances(
    measurements: dict[tuple[str, str], dict[str, str]],
) -> dict[str, float]:
    """Return a stable representative epicentral distance for each station-phase."""
    values: dict[str, list[float]] = {}
    for (station_phase, _), measurement in measurements.items():
        distance = finite_float(measurement.get("epicentral_distance_degrees"))
        if distance is not None:
            values.setdefault(station_phase, []).append(distance)
    return {
        station_phase: sum(distances) / len(distances)
        for station_phase, distances in values.items()
    }


def station_distances(distances: dict[str, float]) -> dict[str, float]:
    """Use one representative distance to keep all phases of a station together."""
    values: dict[str, list[float]] = {}
    for station_phase, distance in distances.items():
        station, _ = station_phase.rsplit(" ", 1)
        values.setdefault(station, []).append(distance)
    return {
        station: sum(station_values) / len(station_values)
        for station, station_values in values.items()
    }


def sort_station_phase_blocks(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
    station_column: int,
    metric_column: int,
    distances: dict[str, float],
    station_distance_values: dict[str, float],
) -> bool:
    """Sort complete metric blocks by distance, then the requested phase order."""
    status_rows = [
        row_number
        for row_number in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(row_number, metric_column).value == "Status"
    ]
    if not status_rows:
        return False
    phase_rank = {phase: index for index, phase in enumerate(PHASE_DISPLAY_ORDER)}
    blocks = []
    for index, start in enumerate(status_rows):
        end = status_rows[index + 1] if index + 1 < len(status_rows) else sheet.max_row + 1
        station_phase = station_phase_for_status_row(
            sheet, start, station_column, metric_column
        )
        station, phase = station_phase.rsplit(" ", 1)
        rows = []
        for row_number in range(start, end):
            rows.append(
                [
                    (sheet.cell(row_number, column).value, copy(sheet.cell(row_number, column)._style))
                    for column in range(1, sheet.max_column + 1)
                ]
            )
        blocks.append((station_phase, station, phase, rows))
    ordered = sorted(
        blocks,
        key=lambda item: (
            station_distance_values.get(item[1], float("inf")),
            phase_rank.get(item[2], len(phase_rank)),
            item[1],
        ),
    )
    if [block[0] for block in ordered] == [block[0] for block in blocks]:
        return False
    target_row = header_row + 1
    for _, _, _, rows in ordered:
        for source_row in rows:
            for column, (value, style) in enumerate(source_row, start=1):
                target = sheet.cell(target_row, column)
                target._style = copy(style)
                target.value = value
            target_row += 1
    return True


def write_station_phase_metadata(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    station_column: int,
    metric_rows: dict[tuple[str, str], int],
    distances: dict[str, float],
) -> int:
    """Display station, phase, and distance on the first three rows of a block."""
    changed = 0
    for (station_phase, metric), row_number in metric_rows.items():
        station, phase = station_phase.rsplit(" ", 1)
        expected: str | None
        if metric == "Status":
            expected = station
        elif metric == "Differential time shift (s)":
            expected = phase
        elif metric == "Correlation":
            distance = distances.get(station_phase)
            expected = f"{distance:.2f}°" if distance is not None else None
        else:
            expected = None
        cell = sheet.cell(row_number, station_column)
        if cell.value != expected:
            changed += 1
        cell.value = expected
    return changed


def refresh_station_acceptance_matrix(
    matrix_path: Path,
    measurements_path: Path,
    output_copy: Path,
) -> dict[str, int]:
    measurements = read_measurements(measurements_path)
    workbook = openpyxl.load_workbook(matrix_path)
    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise MatrixUpdateError(f"Workbook lacks required {SHEET_NAME!r} sheet")
    sheet = workbook[SHEET_NAME]
    header_row = find_header_row(sheet)
    headers = {
        str(cell.value).strip(): cell.column
        for cell in sheet[header_row]
        if cell.value is not None and str(cell.value).strip()
    }
    if "Station" not in headers or "Metric" not in headers:
        workbook.close()
        raise MatrixUpdateError("Station acceptance matrix headers are malformed")
    station_column = headers["Station"]
    metric_column = headers["Metric"]
    pair_columns = {
        label: column
        for label, column in headers.items()
        if label not in {"Station", "Metric"}
    }
    added_pair_columns = add_missing_pair_columns(
        sheet, header_row, pair_columns, measurements
    )
    obsolete_phase_blocks_removed = remove_obsolete_phase_blocks(
        sheet,
        header_row,
        station_column,
        metric_column,
        measurements,
        {"PcP", "PKiKP"},
    )
    added_station_phase_blocks = add_missing_station_phase_blocks(
        sheet,
        header_row,
        station_column,
        metric_column,
        pair_columns,
        measurements,
    )
    conditional_format_ranges_extended = extend_phase_conditional_formats(
        sheet,
        header_row,
        pair_columns,
    )
    distances = phase_distances(measurements)
    station_blocks_sorted = sort_station_phase_blocks(
        sheet,
        header_row,
        station_column,
        metric_column,
        distances,
        station_distances(distances),
    )

    metric_rows: dict[tuple[str, str], int] = {}
    manual_values_before: dict[tuple[str, str], Any] = {}
    manual_statuses_before: dict[tuple[str, str], str] = {}
    current_station_phase: str | None = None
    for row_number in range(header_row + 1, sheet.max_row + 1):
        station_phase_value = sheet.cell(row_number, station_column).value
        metric = sheet.cell(row_number, metric_column).value
        if metric == "Status" and station_phase_value is not None and str(station_phase_value).strip():
            current_station_phase = station_phase_for_status_row(
                sheet, row_number, station_column, metric_column
            )
        if current_station_phase is None or metric is None:
            continue
        key = (current_station_phase, str(metric).strip())
        if key in metric_rows:
            workbook.close()
            raise MatrixUpdateError(f"Duplicate matrix row for {key[0]} / {key[1]}")
        metric_rows[key] = row_number
        if key[1] == MANUAL_METRIC:
            for pair, column in pair_columns.items():
                if (
                    pair in added_pair_columns
                    and (key[0], pair) in measurements
                    and sheet.cell(row_number, column).value is None
                ):
                    sheet.cell(row_number, column).value = 0.0
                phase = key[0].rsplit(" ", 1)[-1]
                if (
                    phase in {"PcP", "ScP"}
                    and (key[0], pair) not in measurements
                    and sheet.cell(row_number, column).value == 0
                ):
                    # Newly added PcP/ScP rows use zero only for an existing
                    # trace; remove harmless placeholder zeros for absent data.
                    sheet.cell(row_number, column).value = None
                manual_values_before[(key[0], pair)] = sheet.cell(
                    row_number, column
                ).value
        elif key[1] == "Status":
            for pair, column in pair_columns.items():
                status = str(sheet.cell(row_number, column).value).strip().upper()
                if status in {"X", "B"}:
                    manual_statuses_before[(key[0], pair)] = status

    for (station_phase, metric), row_number in metric_rows.items():
        if metric not in CALCULATED_METRICS:
            continue
        for column in pair_columns.values():
            sheet.cell(row_number, column).value = None

    updated_measurements = 0
    for (station_phase, pair), measurement in measurements.items():
        if pair not in pair_columns:
            workbook.close()
            raise MatrixUpdateError(f"Measurement pair {pair} has no matrix column")
        values = measurement_values(measurement)
        for metric, value in values.items():
            row_number = metric_rows.get((station_phase, metric))
            if row_number is None:
                workbook.close()
                raise MatrixUpdateError(
                    f"Measurement {station_phase} has no {metric!r} matrix row"
                )
            sheet.cell(row_number, pair_columns[pair]).value = value
        updated_measurements += 1

    for (station_phase, pair), status in manual_statuses_before.items():
        status_row = metric_rows.get((station_phase, "Status"))
        if status_row is None or pair not in pair_columns:
            workbook.close()
            raise MatrixUpdateError(
                f"Stored exclusion {station_phase} {pair} has no Status cell"
            )
        sheet.cell(status_row, pair_columns[pair]).value = status

    nan_format_rule_added = ensure_uncomputed_correlation_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    exclusion_format_rule_added = ensure_exclusion_highlight_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    manual_background_cells_cleared = clear_manual_row_backgrounds(
        sheet,
        metric_rows,
    )
    accepted_blue_rule_added = ensure_accepted_status_blue_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    b_status_rule_added = ensure_b_status_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    manual_orange_rule_added = ensure_nonzero_manual_shift_format(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    passing_format_rules_added = ensure_passing_threshold_formats(
        sheet,
        header_row,
        metric_column,
        pair_columns,
    )
    station_labels_changed = write_station_phase_metadata(
        sheet,
        station_column,
        metric_rows,
        distances,
    )

    source_label = f"Source run: {measurements_path.parent.name}"
    sheet.cell(3, 1).value = source_label
    sheet.cell(4, 1).value = (
        "All station-phase rows from the 16-pair run are retained in distance "
        "order. Enter a manual pick-alignment shift in its labeled row; 0.000 "
        "means no manual adjustment. X marks redundant or unpickable data; B "
        "marks a well-defined shift judged unlikely to result from differential "
        "location. Both are measured and plotted but excluded from the location "
        "fit. X calculated cells are yellow and B status cells are purple. "
        "Accepted A status cells are blue, and non-zero manual pick-alignment "
        "shift cells are orange. "
        "Green numeric cells meet the display thresholds: absolute shift at most "
        "0.05 s, correlation at least 0.85, and minimum SNR at least 1.0. Blank "
        "cells have no measurement; NaN means correlation was not computed."
    )
    for (station_phase, pair), prior_value in manual_values_before.items():
        row_number = metric_rows[(station_phase, MANUAL_METRIC)]
        if sheet.cell(row_number, pair_columns[pair]).value != prior_value:
            workbook.close()
            raise MatrixUpdateError(
                f"Manual shift changed unexpectedly for {station_phase} {pair}"
            )

    temporary_path = matrix_path.with_name(f".{matrix_path.name}.tmp")
    try:
        workbook.save(temporary_path)
    finally:
        workbook.close()
    os.replace(temporary_path, matrix_path)
    output_copy.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(matrix_path, output_copy)
    return {
        "measurement_rows_updated": updated_measurements,
        "manual_cells_preserved": len(manual_values_before),
        "nan_format_rule_added": int(nan_format_rule_added),
        "exclusion_format_rule_added": int(exclusion_format_rule_added),
        "manual_background_cells_cleared": manual_background_cells_cleared,
        "accepted_blue_rule_added": int(accepted_blue_rule_added),
        "b_status_rule_added": int(b_status_rule_added),
        "manual_orange_rule_added": int(manual_orange_rule_added),
        "passing_format_rules_added": passing_format_rules_added,
        "station_labels_changed": station_labels_changed,
        "exclusions_preserved": len(manual_statuses_before),
        "manual_statuses_preserved": len(manual_statuses_before),
        "station_phase_blocks_added": added_station_phase_blocks,
        "obsolete_phase_blocks_removed": obsolete_phase_blocks_removed,
        "conditional_format_ranges_extended": conditional_format_ranges_extended,
        "station_blocks_sorted": int(station_blocks_sorted),
        "pair_columns_added": len(added_pair_columns),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--measurements", type=Path, required=True)
    parser.add_argument("--output-copy", type=Path, required=True)
    args = parser.parse_args()
    counts = refresh_station_acceptance_matrix(
        args.matrix.resolve(),
        args.measurements.resolve(),
        args.output_copy.resolve(),
    )
    print(f"updated={counts['measurement_rows_updated']}")
    print(f"manual_cells_preserved={counts['manual_cells_preserved']}")
    print(args.matrix.resolve())
    print(args.output_copy.resolve())


if __name__ == "__main__":
    main()
