#!/usr/bin/env python3
"""Reorder pair columns in the station-acceptance workbook without changing data."""

from __future__ import annotations

import argparse
import os
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell


SHEET_NAME = "Station status"


def header_row(sheet: openpyxl.worksheet.worksheet.Worksheet) -> int:
    for row in sheet.iter_rows():
        labels = {str(cell.value).strip() for cell in row if cell.value is not None}
        if {"Station", "Metric"}.issubset(labels):
            return row[0].row
    raise ValueError("Station acceptance sheet has no Station/Metric header row")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("pairs", nargs="+", help="Complete desired pair-column order")
    args = parser.parse_args()

    path = args.workbook.resolve()
    workbook = openpyxl.load_workbook(path)
    try:
        sheet = workbook[SHEET_NAME]
        row = header_row(sheet)
        columns = {
            str(cell.value).strip(): cell.column
            for cell in sheet[row]
            if cell.value is not None and str(cell.value).strip()
        }
        current_pairs = [
            str(sheet.cell(row, column).value).strip()
            for column in range(columns["Metric"] + 1, sheet.max_column + 1)
            if sheet.cell(row, column).value is not None
        ]
        if set(current_pairs) != set(args.pairs) or len(current_pairs) != len(args.pairs):
            raise ValueError(
                "Requested pair columns do not exactly match workbook columns: "
                f"current={current_pairs}, requested={args.pairs}"
            )
        first_pair_column = columns["Metric"] + 1
        source_columns = {pair: columns[pair] for pair in current_pairs}
        snapshots = {}
        for pair, source_column in source_columns.items():
            snapshots[pair] = (
                sheet.column_dimensions[
                    openpyxl.utils.get_column_letter(source_column)
                ].width,
                [
                    (
                        sheet.cell(row_number, source_column).value,
                        copy(sheet.cell(row_number, source_column)._style),
                        sheet.cell(row_number, source_column).number_format,
                        copy(sheet.cell(row_number, source_column).alignment),
                        copy(sheet.cell(row_number, source_column).protection),
                    )
                    for row_number in range(1, sheet.max_row + 1)
                ],
            )
        for offset, pair in enumerate(args.pairs):
            target_column = first_pair_column + offset
            source_width, copied = snapshots[pair]
            for row_number, (value, style, number_format, alignment, protection) in enumerate(
                copied, start=1
            ):
                target = sheet.cell(row_number, target_column)
                if isinstance(target, MergedCell):
                    continue
                target.value = value
                target._style = style
                target.number_format = number_format
                target.alignment = alignment
                target.protection = protection
            sheet.column_dimensions[
                openpyxl.utils.get_column_letter(target_column)
            ].width = source_width
        final_pairs = [
            str(sheet.cell(row, column).value).strip()
            for column in range(first_pair_column, first_pair_column + len(args.pairs))
        ]
        if final_pairs != args.pairs:
            raise ValueError(f"Pair-column verification failed: {final_pairs}")
        temporary = path.with_name(f".{path.name}.tmp")
        workbook.save(temporary)
    finally:
        workbook.close()
    os.replace(temporary, path)
    print(", ".join(args.pairs))


if __name__ == "__main__":
    main()
