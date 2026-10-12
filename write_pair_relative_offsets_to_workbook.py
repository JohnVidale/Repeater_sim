#!/usr/bin/env python3
"""Write preferred event-2-minus-event-1 offsets to the working workbook."""

from __future__ import annotations

import argparse
import csv
import shutil
import math
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from aligned_pairs import aligned_rows, header_map, refresh_reference_times


DEFAULT_WORKBOOK = Path("ICevents_working.xlsx")
DEFAULT_RELATIVE_LOCATIONS = Path(
    "outputs/multiphase_median_cc0.5_workbook_20260813_165113/"
    "median_absolute_relative_locations_no_pkikp.csv"
)


def normalized_header(row: tuple[Any, ...]) -> dict[str, int]:
    return {
        str(value).strip().lower(): index + 1
        for index, value in enumerate(row)
        if value is not None and str(value).strip()
    }


def read_relative_locations(path: Path) -> dict[str, dict[str, float | str | int]]:
    rows: dict[str, dict[str, float | str | int]] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            pair = row["pair"].strip()
            # A pair with fewer than four accepted measurements is retained in
            # the fit table with intentionally blank location fields.  It has
            # no defensible offset to write to the workbook.
            if not str(row.get("event1_lat", "")).strip():
                continue
            rows[pair] = {
                "phase_set": row["phase_set"],
                "n": int(row["n"]),
                "event1": int(row["event1"]),
                "event2": int(row["event2"]),
                "event1_lat": float(row["event1_lat"]),
                "event1_lon": float(row["event1_lon"]),
                "event1_depth_km": float(row["event1_depth_km"]),
                "event2_lat": float(row["event2_lat"]),
                "event2_lon": float(row["event2_lon"]),
                "event2_depth_km": float(row["event2_depth_km"]),
                "delta_lat_degrees": float(row["delta_lat_degrees"]),
                "delta_lon_degrees": float(row["delta_lon_degrees"]),
                "event2_minus_event1_origin_time_shift_s": float(
                    row["event2_minus_event1_origin_time_shift_s"]
                ),
                "east_km": float(row["east_km"]),
                "north_km": float(row["north_km"]),
                "horizontal_km": float(row["horizontal_km"]),
                "separation_3d_km": float(row["separation_3d_km"]),
                "median_abs_fit_residual_s": float(row["median_abs_fit_residual_s"]),
                "residual_rms_s": float(row["residual_rms_s"]),
            }
    return rows


def ensure_columns(sheet, names: list[str]) -> dict[str, int]:
    header = normalized_header(
        tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    )
    next_col = sheet.max_column + 1
    for name in names:
        key = name.lower()
        if key in header:
            continue
        sheet.cell(row=1, column=next_col).value = name
        header[key] = next_col
        next_col += 1
    return header


def ensure_column_after(sheet, name: str, after_name: str) -> dict[str, int]:
    """Ensure a result column directly follows its catalog coordinate."""
    header = normalized_header(
        tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    )
    target_column = header[after_name.lower()] + 1
    existing_column = header.get(name.lower())
    if existing_column is not None:
        if existing_column != target_column:
            raise ValueError(
                f"{name} must directly follow {after_name}; "
                f"found it in column {existing_column}"
            )
        return header
    sheet.insert_cols(target_column)
    sheet.cell(row=1, column=target_column).value = name
    return normalized_header(
        tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    )


def write_offsets(workbook_path: Path, relative_locations_path: Path) -> Path:
    locations = read_relative_locations(relative_locations_path)
    backup_path = (
        workbook_path.parent
        / f"{workbook_path.stem}_before_pair_relative_offsets_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}{workbook_path.suffix}"
    )
    shutil.copy2(workbook_path, backup_path)

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook["aligned_pairs"]
        header = header_map(sheet)
        written = 0
        for label, row_index in aligned_rows(sheet).items():
            if label not in locations:
                continue
            row = locations[label]
            delta_lat = float(row["delta_lat_degrees"])
            delta_lon = float(row["delta_lon_degrees"])
            delta_time = float(row["event2_minus_event1_origin_time_shift_s"])
            values = {
                "relative_location_phase_set": row["phase_set"],
                "relative_location_n": row["n"],
                "relative_location_north_m": 111195.0 * delta_lat,
                "relative_location_east_m": 111195.0 * math.cos(math.radians((float(row["event1_lat"]) + float(row["event2_lat"])) / 2.0)) * delta_lon,
                "relative_location_event2_minus_event1_shift_s": delta_time,
                "relative_location_3d_m": 1000.0 * float(row["separation_3d_km"]),
                "relative_location_median_abs_residual_s": row[
                    "median_abs_fit_residual_s"
                ],
                "relative_location_rms_residual_s": row["residual_rms_s"],
                "relative_location_event1_lat": row["event1_lat"],
                "relative_location_event1_lon": row["event1_lon"],
                "relative_location_event2_lat": row["event2_lat"],
                "relative_location_event2_lon": row["event2_lon"],
            }
            for name, value in values.items():
                cell = sheet.cell(row=row_index, column=header[name.lower()])
                cell.value = value
                if isinstance(value, float):
                    cell.number_format = "0.0000"
            sheet.cell(row_index, header["relative_location_available"]).value = True
            refresh_reference_times(workbook, row_index)
            written += 1
        workbook.save(workbook_path)
    finally:
        workbook.close()
    print(f"backup={backup_path}")
    print(f"written={written}")
    return backup_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--relative-locations", type=Path, default=DEFAULT_RELATIVE_LOCATIONS)
    args = parser.parse_args()
    write_offsets(args.workbook, args.relative_locations)


if __name__ == "__main__":
    main()
